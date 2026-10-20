import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=False)
if "00_BANG_CHAO_GIA" in wb.sheetnames:
    ws = wb["00_BANG_CHAO_GIA"]
    print(f"00_BANG_CHAO_GIA Max Row: {ws.max_row}")
    for r in range(1, min(30, ws.max_row)):
        print(f"Row {r}, Col 14 (Total): {ws.cell(r, 14).value}")
else:
    print("No 00_BANG_CHAO_GIA sheet")
