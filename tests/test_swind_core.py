"""
Unit test cho các cơ chế kiểm soát của pipeline Swind (chạy: ./venv/bin/python -m unittest discover -s tests -p "test_swind_*.py").
"""
import os
import sys
import json
import tempfile
import unittest

SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src")
sys.path.insert(0, SRC)

import openpyxl  # noqa: E402
from agent_08_qa_checker import Agent08_QAChecker  # noqa: E402
from swind_master_runner import evaluate_quotation_gate  # noqa: E402
from swind_utils import K, prune_named_styles_in_file  # noqa: E402
import agent_02_boq_reviewer as reviewer  # noqa: E402


def _read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _item(row, w, h, price, cat="VAN_GIO", subtype="VCD", qty=1):
    return {"source_sheet": "S", "source_row": row, "raw_name": f"Van {w}x{h}", "category": cat, "qty": qty,
            "w": w, "h": h, "d": 0, "calc_sell_unit": price, "resolved_spec": {"subtype": subtype, "model_std": "X"}}


class ValueAuditTest(unittest.TestCase):
    def setUp(self):
        self.qa = Agent08_QAChecker.__new__(Agent08_QAChecker)

    def test_unpriced_manufactured_item_is_critical(self):
        issues = self.qa.run_value_audit([_item(5, 400, 300, 0)])
        self.assertEqual([i["error_type"] for i in issues], ["UNPRICED_ITEM"])
        self.assertEqual(issues[0]["severity"], "CRITICAL")

    def test_bigger_but_much_cheaper_is_flagged(self):
        items = [_item(1, 300, 200, 500000), _item(2, 800, 600, 200000)]
        issues = self.qa.run_value_audit(items)
        self.assertEqual([(i["error_type"], i["row"]) for i in issues], [("PRICE_OUTLIER", 2)])

    def test_monotonic_prices_pass(self):
        items = [_item(1, 300, 200, 200000), _item(2, 600, 400, 300000), _item(3, 800, 600, 450000)]
        self.assertEqual(self.qa.run_value_audit(items), [])

    def test_no_comparison_across_1000mm_tier(self):
        # Van cạnh >= 1000mm có công thức riêng -> không so với van < 1000mm
        items = [_item(1, 1050, 350, 1050000), _item(2, 850, 550, 665000)]
        self.assertEqual(self.qa.run_value_audit(items), [])

    def test_small_items_with_high_per_m2_price_pass(self):
        # Van nhỏ có chi phí cố định lớn -> đơn giá/m2 cao là bình thường
        items = [_item(1, 150, 100, 160000), _item(2, 1000, 800, 900000)]
        self.assertEqual(self.qa.run_value_audit(items), [])


class GateTest(unittest.TestCase):
    def test_missing_dims_block(self):
        cov = {"S": {"missing_dims": ["Dòng 5"], "skipped_in_scope": [], "suspect_unparsed": []}}
        gate = evaluate_quotation_gate(cov, {}, {}, [])
        self.assertTrue(gate["blocking"])
        self.assertEqual(gate["unpriced_rows"], 1)

    def test_parse_error_blocks(self):
        gate = evaluate_quotation_gate({}, {}, {}, [], parse_errors={"Sheet A": "RuntimeError: x"})
        self.assertTrue(any("Sheet A" in b for b in gate["blocking"]))

    def test_clean_run_passes(self):
        cov = {"S": {"missing_dims": [], "skipped_in_scope": [], "suspect_unparsed": []}}
        gate = evaluate_quotation_gate(cov, {"S": {"issues": []}}, {"critical_errors": []}, [])
        self.assertEqual((gate["blocking"], gate["confirm"]), ([], []))

    def test_assumed_dims_need_confirmation(self):
        it = {"source_row": 9, "raw_name": "Van VCD", "qty": 2, "dim_assumption": "300x300"}
        gate = evaluate_quotation_gate({}, {}, {}, [], all_items=[it])
        self.assertTrue(any("GIẢ ĐỊNH" in c for c in gate["confirm"]))


class FireDamperRuleTest(unittest.TestCase):
    def _resolve(self, name, note=""):
        from agent_01_boq_extractor import Agent01_BOQExtractor
        from swind_utils import get_subtype_rules
        it = {"raw_name": name, "name": name, "note_boq": note, "category": "VAN_GIO", "unit": "cái", "qty": 1,
              "w": 600.0, "h": 400.0, "d": 0.0, "l": 0.0, "thickness": 0.0, "model": "", "ei_rating": 0}
        Agent01_BOQExtractor.enrich_all_items([it], get_subtype_rules())
        return it

    def test_fd_without_ei_defaults_to_ei120(self):
        for nm in ("Van chặn lửa FD 600x400", "Van chặn lửa cầu chì 600x400", "Van FD EI45 600x400"):
            it = self._resolve(nm)
            self.assertIn("2.01", it["resolved_spec"]["model_std"] or it["resolved_spec"]["model_fslf"], nm)
            self.assertFalse(it.get("fd_non_ei_keyword"))

    def test_explicit_plain_fd_is_flagged(self):
        it = self._resolve("Van chặn lửa FD thường 600x400")
        self.assertNotIn("2.01", it["resolved_spec"]["model_std"] or it["resolved_spec"]["model_fslf"])
        self.assertEqual(it.get("fd_non_ei_keyword"), "fd thường")
        gate = evaluate_quotation_gate({}, {}, {}, [], all_items=[dict(it, source_row=7)])
        self.assertTrue(any("HỎI LẠI USER" in c for c in gate["confirm"]))


class PricingConstantsTest(unittest.TestCase):
    def test_lookup_and_fallback(self):
        self.assertEqual(K("van.ton_m2.bac_lon", -1), 321300.0)
        self.assertEqual(K("khong.ton.tai", 123.0), 123.0)


class PruneStylesTest(unittest.TestCase):
    def test_prune_keeps_cell_formatting(self):
        from openpyxl.styles import NamedStyle, Font
        with tempfile.TemporaryDirectory() as td:
            src, dst = os.path.join(td, "a.xlsx"), os.path.join(td, "b.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            for i in range(300):
                wb.add_named_style(NamedStyle(name=f"junk{i}"))
            used = NamedStyle(name="used", font=Font(bold=True, color="FF0000"))
            wb.add_named_style(used)
            ws["A1"] = "x"
            ws["A1"].style = "used"
            ws["B2"] = 5
            wb.save(src)
            removed = prune_named_styles_in_file(src, dst)
            self.assertGreaterEqual(removed, 300)
            c = openpyxl.load_workbook(dst)["Sheet"]["A1"]
            self.assertTrue(c.font.b)
            self.assertEqual(c.font.color.rgb, "00FF0000")
            self.assertEqual(c.value, "x")


class LearningTest(unittest.TestCase):
    def test_auto_suggestion_goes_to_queue_not_live_dict(self):
        with tempfile.TemporaryDirectory() as td:
            orig = (reviewer.LEARNING_PATH, reviewer.QUEUE_PATH, reviewer.resolve_path)
            reviewer.resolve_path = lambda rel: os.path.join(td, os.path.basename(rel))
            try:
                reviewer.save_learned_keyword("hộp xyz", "THUONG_MAI", approved=False, score=0.7)
                self.assertFalse(os.path.exists(os.path.join(td, "boq_keyword_learning.json")))
                q = _read_json(os.path.join(td, "pending_approval_queue.json"))
                self.assertEqual(q["pending_proposals"][0]["status"], "PENDING_APPROVAL")
                reviewer.approve_pending_keywords(approve_all=True)
                live = _read_json(os.path.join(td, "boq_keyword_learning.json"))
                self.assertEqual(live["hộp xyz"], "THUONG_MAI")
            finally:
                reviewer.LEARNING_PATH, reviewer.QUEUE_PATH, reviewer.resolve_path = orig


if __name__ == "__main__":
    unittest.main()
