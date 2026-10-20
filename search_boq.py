import openpyxl
import sys

wb = openpyxl.load_workbook("input/Haus_Dalat_BOQ_MEP (NCC).xlsx", data_only=True)
ws = wb["DHKK"]

for r in range(1, ws.max_row + 1):
    val = str(ws.cell(r, 2).value)
    if "0.95" in val or "0,95" in val:
        print(f"Row {r}: {val}")

