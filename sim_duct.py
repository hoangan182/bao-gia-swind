import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=True)
ws01 = wb["01_BOC_TACH_ONG"]

for r in range(2, ws01.max_row + 1):
    comp_name = ws01.cell(r, 3).value
    mat_id = ws01.cell(r, 5).value
    qty = ws01.cell(r, 7).value
    price = ws01.cell(r, 10).value
    if comp_name and "1.15" in comp_name:
        print(f"Row {r}: {comp_name} ({mat_id}) - Qty: {qty} - Price: {price}")
        parent_id = ws01.cell(r, 2).value
        # Find all other items with same parent_id
        print(f"  Parent: {parent_id}")
        for rr in range(2, ws01.max_row + 1):
            if ws01.cell(rr, 2).value == parent_id:
                print(f"    - {ws01.cell(rr, 3).value}: qty={ws01.cell(rr, 7).value}, price={ws01.cell(rr, 10).value}, total={ws01.cell(rr, 11).value}")
        break
