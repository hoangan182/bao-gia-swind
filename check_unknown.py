import openpyxl
import glob

latest_file = sorted(glob.glob("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_*.xlsx"))[-1]
wb = openpyxl.load_workbook(latest_file, data_only=True)
ws = wb["00_TONG_HOP_BAO_GIA"]

unknowns = []
for r in range(2, ws.max_row + 1):
    cat = ws.cell(r, 6).value
    if cat == "UNKNOWN":
        name = ws.cell(r, 4).value
        unknowns.append(f"Row {r}: {name}")

print(f"Có {len(unknowns)} hạng mục không thể phân loại (UNKNOWN).")
if unknowns:
    for u in unknowns[:10]:
        print(u)
