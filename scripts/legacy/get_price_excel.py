import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=True)
ws06 = wb["06_BANG_GIA_VAT_TU_AP_DUNG"]
for r in range(2, ws06.max_row + 1):
    mat_id = ws06.cell(r, 2).value
    price = ws06.cell(r, 5).value
    if mat_id in ["MAT-STEEL-PN-1.15", "MAT-BOX-1.15"]:
        print(f"{mat_id}: {price}")
