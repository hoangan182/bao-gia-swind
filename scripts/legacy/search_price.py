import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=False)
found = False
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None and isinstance(val, (int, float)):
                if abs(val - 273549) < 1 or abs(val - 273.549) < 0.1:
                    print(f"Found {val} in {sheet_name} at R{r}C{c}")
                    found = True
            elif isinstance(val, str) and "273549" in val:
                print(f"Found str {val} in {sheet_name} at R{r}C{c}")
                found = True
if not found:
    print("Not found")
