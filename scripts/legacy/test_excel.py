import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=True)
ws06 = wb["06_BANG_GIA_VAT_TU_AP_DUNG"]
print("Sheet 06 Sample (Rows 2-5):")
for r in range(2, 6):
    print([ws06.cell(r, c).value for c in range(1, 6)])
    
ws00 = wb["00_TONG_HOP_BAO_GIA"]
print("\nSheet 00 Sample (Rows 2-5):")
for r in range(2, 6):
    print([ws00.cell(r, c).value for c in [2,4,7,8,11,18]]) # ID, Name, Qty, Cost_Ong, Tong_VT, ThanhTien
