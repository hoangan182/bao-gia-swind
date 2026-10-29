import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=False)
ws = wb["00_BANG_CHAO_GIA"]
max_row = ws.max_row
print(f"Max Row: {max_row}")
for r in range(max_row - 20, max_row + 1):
    print(f"Row {r}, Col 14 (Total): {ws.cell(r, 14).value}")
