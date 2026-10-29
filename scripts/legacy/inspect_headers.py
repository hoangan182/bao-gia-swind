import openpyxl

wb = openpyxl.load_workbook("input/Haus_Dalat_BOQ_MEP (NCC).xlsx", data_only=True)
ws = wb["DHKK"]

for r in range(1, 15):
    row_vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v: row_vals.append(f"Col {c}: {str(v).strip()}")
    if row_vals:
        print(f"Row {r}:", " | ".join(row_vals))
