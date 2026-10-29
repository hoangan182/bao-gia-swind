import openpyxl

wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=False)
ws06 = wb["06_BANG_GIA_VAT_TU_AP_DUNG"]

prices = {}
for r in range(2, ws06.max_row + 1):
    mat_id = ws06.cell(r, 2).value
    price = ws06.cell(r, 5).value
    if mat_id:
        prices[mat_id] = price

ws01 = wb["01_BOC_TACH_ONG"]
total_ong = 0
for r in range(2, ws01.max_row + 1):
    mat_id = ws01.cell(r, 5).value
    if not mat_id: continue
    qty = ws01.cell(r, 9).value  # wait, I is a formula! I = G * H. G = calc_qty_per_item, H = parent_qty
    # let's just get G
    qty_per_item = ws01.cell(r, 7).value or 0
    # parent qty is from 00, let's assume 1 for now to see if price is fetched
    price = prices.get(mat_id, 0)
    if price > 0:
        total_ong += qty_per_item * price

print(f"Total calculated cost in 01_BOC_TACH_ONG (assuming parent qty=1): {total_ong}")
