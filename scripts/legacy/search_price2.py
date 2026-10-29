import openpyxl
wb = openpyxl.load_workbook("input/Haus_Dalat_BOQ_MEP (NCC).xlsx", data_only=True)
ws = wb["DHKK"]
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        val = ws.cell(r, c).value
        if val is not None and isinstance(val, (int, float)):
            if abs(val - 273549) < 1:
                print(f"Found {val} in original at R{r}C{c}")
                name = ws.cell(r, 3).value
                print(f"Name: {name}")
