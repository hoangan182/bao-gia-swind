import openpyxl
wb = openpyxl.load_workbook("input/Haus_Dalat_BOQ_MEP (NCC).xlsx", data_only=True)
ws = wb["DHKK"]
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        val = ws.cell(r, c).value
        if isinstance(val, str) and "1.15" in val and "Z180" in val:
            print(f"Found in original at R{r}C{c}: {val}")
