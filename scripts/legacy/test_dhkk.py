import openpyxl
wb = openpyxl.load_workbook("input/Haus_Dalat_BOQ_MEP (NCC).xlsx", data_only=False)
ws = wb["DHKK"]
for r in range(ws.max_row - 20, ws.max_row + 1):
    val = ws.cell(r, 16).value # usually total is at 16 or somewhere
    val14 = ws.cell(r, 14).value
    print(f"Row {r} Col 14: {val14}, Col 16: {val}")
