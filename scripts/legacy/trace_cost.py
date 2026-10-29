import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=True)
ws00 = wb["00_TONG_HOP_BAO_GIA"]
ws01 = wb["01_BOC_TACH_ONG"]

item_id = None
item_name = None
for r in range(2, ws00.max_row + 1):
    name = ws00.cell(r, 4).value
    if name and "1.15mm phủ lớp mạ kẽm Z180" in name:
        item_id = ws00.cell(r, 2).value
        item_name = name
        print(f"Found item in Sheet 00: {item_id} - {item_name}")
        print(f"  Col 11 (Vật tư): {ws00.cell(r, 11).value}")
        print(f"  Col 12 (Nhân công): {ws00.cell(r, 12).value}")
        print(f"  Col 13 (Xưởng): {ws00.cell(r, 13).value}")
        print(f"  Col 14 (Tổng vốn): {ws00.cell(r, 14).value}")
        print(f"  Col 15 (Vốn/ĐVT): {ws00.cell(r, 15).value}")
        print(f"  Col 16 (Margin): {ws00.cell(r, 16).value}")
        print(f"  Col 17 (Giá bán): {ws00.cell(r, 17).value}")
        break

if item_id:
    print(f"\nComponents in 01_BOC_TACH_ONG for {item_id}:")
    for r in range(2, ws01.max_row + 1):
        if ws01.cell(r, 2).value == item_id:
            comp = ws01.cell(r, 3).value
            qty = ws01.cell(r, 7).value
            unit_price = ws01.cell(r, 10).value
            cost = qty * unit_price if qty and unit_price else 0
            print(f"  - {comp}: qty={qty}, unit_price={unit_price} => cost={cost}")
