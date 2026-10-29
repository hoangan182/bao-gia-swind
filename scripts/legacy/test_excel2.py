import openpyxl
wb = openpyxl.load_workbook("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_20261011_100917.xlsx", data_only=False)
ws00 = wb["00_TONG_HOP_BAO_GIA"]
max_row = ws00.max_row
print(f"Sheet 00 max row: {max_row}")
print(f"Total formula at bottom: {ws00.cell(max_row, 18).value}")
for r in range(2, max_row):
    if ws00.cell(r, 2).value and "SW-DUCT" in ws00.cell(r, 2).value:
        print(f"Found Duct at row {r}. ID: {ws00.cell(r,2).value}. Formula for Thanh Tien (Col 18): {ws00.cell(r,18).value}")
        print(f"Formula for Col 11 (Tong VT): {ws00.cell(r,11).value}")
        break

ws01 = wb["01_BOC_TACH_ONG"]
for r in range(2, ws01.max_row + 1):
    if ws01.cell(r, 2).value and "SW-DUCT" in ws01.cell(r, 2).value:
        print(f"\nSheet 01 - Found Duct at row {r}.")
        print(f"Mat ID (Col 5): {ws01.cell(r,5).value}")
        print(f"Formula for Unit Price (Col 10): {ws01.cell(r,10).value}")
        break
