import os
import sys
import glob
import json
import openpyxl
from datetime import datetime


def resolve_path(rel_path):
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, rel_path)


COMMERCIAL_TECH_PATH = resolve_path("data/master_data/master_commercial_tech.json")
PRICES_DIR = resolve_path("data/master_data/prices")


def get_latest_commercial_price_file():
    if not os.path.exists(PRICES_DIR):
        os.makedirs(PRICES_DIR, exist_ok=True)
    pattern = os.path.join(PRICES_DIR, "master_commercial_price_*.json")
    files = glob.glob(pattern)
    if not files:
        return resolve_path("data/master_data/master_commercial_price.json")
    files.sort()
    return files[-1]


def run_learner(excel_file):
    print("=" * 60)
    print(f"[Agent 09 - Price Learner] Đang đọc file: {os.path.basename(excel_file)}")

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        if "06_BANG_GIA_VAT_TU_AP_DUNG" not in wb.sheetnames:
            print("[LỖI] Không tìm thấy Sheet 06_BANG_GIA_VAT_TU_AP_DUNG trong file này.")
            return

        ws = wb["06_BANG_GIA_VAT_TU_AP_DUNG"]

        learned_items = []
        for r in range(2, ws.max_row + 1):
            mat_id = ws.cell(r, 2).value
            name = ws.cell(r, 3).value
            category = ws.cell(r, 4).value
            price = ws.cell(r, 5).value

            if not mat_id:
                continue

            if str(category).strip() in ("THUONG_MAI", "ONG_MEM", "KHAC") or any(p in str(mat_id) for p in ("MAT-KHAC", "MAT-THUONG_MAI", "MAT-ONG_MEM", "MAT-BOX", "MAT-FILTER", "MAT-LUOI")):
                try:
                    price = float(price) if price else 0.0
                except ValueError:
                    price = 0.0

                if price > 0:
                    learned_items.append({
                        "MAT_ID": mat_id,
                        "Name": name,
                        "Price": price
                    })

        if not learned_items:
            print(" -> Không tìm thấy mặt hàng thương mại nào có đơn giá > 0 để học.")
            return

        print(f" -> Đã trích xuất {len(learned_items)} mặt hàng thương mại có giá.")

        # 1. Update Tech DB
        with open(COMMERCIAL_TECH_PATH, "r", encoding="utf-8") as f:
            tech_data = json.load(f)

        existing_tech_ids = {item["MAT_ID"]: i for i, item in enumerate(tech_data)}
        for item in learned_items:
            if item["MAT_ID"] in existing_tech_ids:
                tech_data[existing_tech_ids[item["MAT_ID"]]]["Name"] = item["Name"]
            else:
                tech_data.append({
                    "MAT_ID": item["MAT_ID"],
                    "Category": "THƯƠNG_MẠI_TỰ_ĐỘNG",
                    "Name": item["Name"],
                    "Unit": "Cái/Bộ",
                    "Source_Trace": "Auto-learned from Excel"
                })

        with open(COMMERCIAL_TECH_PATH, "w", encoding="utf-8") as f:
            json.dump(tech_data, f, indent=4, ensure_ascii=False)

        # 2. Update Price DB (Append-only Audit Mechanism)
        latest_price_file = get_latest_commercial_price_file()
        with open(latest_price_file, "r", encoding="utf-8") as f:
            price_data = json.load(f)

        existing_price_ids = {item["MAT_ID"]: i for i, item in enumerate(price_data)}
        for item in learned_items:
            if item["MAT_ID"] in existing_price_ids:
                price_data[existing_price_ids[item["MAT_ID"]]]["Price"] = item["Price"]
            else:
                price_data.append({
                    "MAT_ID": item["MAT_ID"],
                    "Price": item["Price"]
                })

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        new_price_file = os.path.join(PRICES_DIR, f"master_commercial_price_{timestamp}.json")

        with open(new_price_file, "w", encoding="utf-8") as f:
            json.dump(price_data, f, indent=4, ensure_ascii=False)

        print(f"[THÀNH CÔNG] [Agent 09] Đã học và cập nhật {len(learned_items)} mức giá mới!")
        print(f" -> File gốc (không bị ghi đè): {os.path.basename(latest_price_file)}")
        print(f" -> Sinh file Audit mới tại: {new_price_file}")

    except Exception as e:
        print(f"[LỖI] Đã xảy ra lỗi trong quá trình học giá: {e}")


class Agent09_PriceLearner:
    """
    Agent 09 - Commercial Price Learner (Học giá vật tư mới từ Sheet 06 với cơ chế Append-only Audit)
    """
    run_learner = staticmethod(run_learner)
    get_latest_commercial_price_file = staticmethod(get_latest_commercial_price_file)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python agent_09_price_learner.py <path_to_excel_file>")
        sys.exit(1)

    excel_path = sys.argv[1]
    if not os.path.exists(excel_path):
        print(f"[LỖI] Không tìm thấy file: {excel_path}")
        sys.exit(1)

    run_learner(excel_path)
