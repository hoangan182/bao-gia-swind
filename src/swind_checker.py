import openpyxl
import json

class SwindChecker:
    """
    Agent 2 (The Checker): Kiểm toán viên độc lập 3 tầng.
    - Tầng 1: Kiểm toán toàn vẹn mã vật tư (Zero missing MAT_ID).
    - Tầng 2: Kiểm toán thứ nguyên kỹ thuật & định mức cơ khí.
    - Tầng 3: Kiểm toán cấu trúc công thức và liên kết bảng.
    """
    def __init__(self, excel_filepath, wb=None):
        self.filepath = excel_filepath
        self.wb = wb if wb is not None else openpyxl.load_workbook(excel_filepath, data_only=False) # Load formula
        self.issues = []

    def run_full_audit(self):
        self.issues = []
        
        # 1. Thu thập danh sách MAT_ID hợp lệ từ Master Data (Sheet 06 - Giá)
        ws_master = self.wb["06_BANG_GIA_VAT_TU_AP_DUNG"]
        valid_mat_ids = {}
        for r in range(2, ws_master.max_row + 1):
            mid = ws_master.cell(r, 2).value
            price = ws_master.cell(r, 5).value or 0
            if mid:
                valid_mat_ids[str(mid).strip()] = price
                
        # 2. Thu thập danh sách PARENT_ITEM_ID từ Bảng tổng hợp
        ws_summary = self.wb["00_TONG_HOP_BAO_GIA"]
        valid_parent_ids = set()
        for r in range(2, ws_summary.max_row + 1):
            pid = ws_summary.cell(r, 2).value
            if pid:
                valid_parent_ids.add(str(pid).strip())

        # 3. Kiểm toán các sheet con (01_ONG, 02_CUA, 03_VAN)
        sub_sheets = [
            ("01_BOC_TACH_ONG", "ỐNG GIÓ"),
            ("02_BOC_TACH_CUA", "CỬA GIÓ"),
            ("03_BOC_TACH_VAN", "VAN GIÓ")
        ]
        
        total_components_audited = 0
        
        for sheet_name, sheet_label in sub_sheets:
            ws = self.wb[sheet_name]
            for r in range(2, ws.max_row + 1):
                sub_key = ws.cell(r, 1).value
                parent_id = ws.cell(r, 2).value
                comp_name = ws.cell(r, 3).value
                mat_id = ws.cell(r, 5).value
                unit = ws.cell(r, 6).value
                qty_per_item = ws.cell(r, 7).value
                
                if not sub_key and not parent_id:
                    continue
                    
                total_components_audited += 1
                
                # Check 1.1: Parent ID tồn tại
                if not parent_id or str(parent_id).strip() not in valid_parent_ids:
                    self.issues.append({
                        "sheet": sheet_name,
                        "row": r,
                        "sub_key": str(sub_key),
                        "severity": "CRITICAL",
                        "error_type": "ORPHAN_COMPONENT",
                        "message": f"PARENT_ITEM_ID '{parent_id}' không tồn tại trong Bảng Tổng Hợp 00!"
                    })
                    
                # Check 1.2: MAT_ID tồn tại và có giá > 0 trong Master Data
                if not mat_id or str(mat_id).strip() not in valid_mat_ids:
                    self.issues.append({
                        "sheet": sheet_name,
                        "row": r,
                        "sub_key": str(sub_key),
                        "severity": "CRITICAL",
                        "error_type": "INVALID_MAT_ID",
                        "message": f"Mã vật tư '{mat_id}' không tồn tại trong Bảng Giá 06! Tra cứu giá sẽ lỗi."
                    })
                elif valid_mat_ids.get(str(mat_id).strip(), 0) == 0:
                    self.issues.append({
                        "sheet": sheet_name,
                        "row": r,
                        "sub_key": str(sub_key),
                        "severity": "WARNING",
                        "error_type": "ZERO_PRICE",
                        "message": f"Vật tư '{mat_id}' đang bị thiếu Đơn Giá (Bằng 0) trong Bảng Giá 06! Cần điền giá và chạy tool Học Giá."
                    })
                    
                # Check 2.1: Định mức > 0
                if qty_per_item is None or not isinstance(qty_per_item, (int, float)) or qty_per_item <= 0:
                    self.issues.append({
                        "sheet": sheet_name,
                        "row": r,
                        "sub_key": str(sub_key),
                        "severity": "CRITICAL",
                        "error_type": "ZERO_QTY_PER_ITEM",
                        "message": f"Định mức cấu kiện '{comp_name}' <= 0 ({qty_per_item})!"
                    })
                    
                # Check 3.1: Công thức thành tiền chuẩn
                formula_total = str(ws.cell(r, 11).value) # Col K
                if not formula_total.startswith("=I") or "*" not in formula_total:
                    self.issues.append({
                        "sheet": sheet_name,
                        "row": r,
                        "sub_key": str(sub_key),
                        "severity": "HIGH",
                        "error_type": "CORRUPTED_FORMULA",
                        "message": f"Công thức thành tiền Col K không đúng chuẩn `=I*J`: '{formula_total}'"
                    })

        # 4. Kiểm toán Sheet Tổng Hợp (00_TONG_HOP_BAO_GIA)
        total_items_audited = 0
        for r in range(2, ws_summary.max_row + 1):
            item_id = ws_summary.cell(r, 2).value
            if not item_id: continue
            
            total_items_audited += 1
            f_cp_ong = str(ws_summary.cell(r, 8).value)
            f_cp_cua = str(ws_summary.cell(r, 9).value)
            f_cp_van = str(ws_summary.cell(r, 10).value)
            
            # Check công thức liên kết SUMIF
            if "SUMIF" not in f_cp_ong or "SUMIF" not in f_cp_cua or "SUMIF" not in f_cp_van:
                self.issues.append({
                    "sheet": "00_TONG_HOP_BAO_GIA",
                    "row": r,
                    "item_id": str(item_id),
                    "severity": "HIGH",
                    "error_type": "MISSING_SUMIF",
                    "message": "Thiếu công thức SUMIF liên kết chi phí vật tư từ các sheet con!"
                })

        # Tổng kết chứng chỉ
        has_blocking = any(issue["severity"] in ("CRITICAL", "HIGH") for issue in self.issues)
        qa_status = "QA_PASSED" if not has_blocking else "QA_REJECTED"
        
        report = {
            "qa_status": qa_status,
            "total_items_audited": total_items_audited,
            "total_components_audited": total_components_audited,
            "total_issues_found": len(self.issues),
            "issues": self.issues
        }
        
        return report

if __name__ == "__main__":
    import glob
    import sys
    try:
        latest_file = sorted(glob.glob("data/projects/Haus_Dalat/Haus_Dalat_ESTIMATION_*.xlsx"))[-1]
    except IndexError:
        print("Không tìm thấy file kết quả nào để kiểm tra.")
        sys.exit(1)
        
    print(f"Đang kiểm tra file: {latest_file}")
    checker = SwindChecker(latest_file)
    report = checker.run_full_audit()
    
    print(f"\n--- BÁO CÁO KIỂM TOÁN (QA STATUS: {report['qa_status']}) ---")
    print(f"- Số lượng hạng mục (Items) đã kiểm tra: {report['total_items_audited']}")
    print(f"- Số lượng cấu kiện (BOM) đã kiểm tra: {report['total_components_audited']}")
    print(f"- Số lỗi phát hiện: {report['total_issues_found']}")
    
    if report['total_issues_found'] > 0:
        print("\nCHI TIẾT LỖI:")
        for idx, issue in enumerate(report['issues']):
            print(f"  {idx+1}. [{issue['severity']}] {issue['error_type']} ở {issue['sheet']} (Dòng {issue['row']}): {issue['message']}")
    else:
        print("\nTuyệt vời! Không phát hiện lỗi cấu trúc hay mã vật tư nào.")
