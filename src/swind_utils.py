import re
import unicodedata
import hashlib
import openpyxl
from openpyxl.utils import get_column_letter

VIETNAMESE_CHAR_MAP = {
    'à': 'a', 'á': 'a', 'ả': 'a', 'ã': 'a', 'ạ': 'a',
    'ă': 'a', 'ắ': 'a', 'ằ': 'a', 'ẳ': 'a', 'ẵ': 'a', 'ặ': 'a',
    'â': 'a', 'ấ': 'a', 'ầ': 'a', 'ẩ': 'a', 'ẫ': 'a', 'ậ': 'a',
    'đ': 'd', 'Đ': 'D',
    'è': 'e', 'é': 'e', 'ẻ': 'e', 'ẽ': 'e', 'ẹ': 'e',
    'ê': 'e', 'ế': 'e', 'ề': 'e', 'ể': 'e', 'ễ': 'e', 'ệ': 'e',
    'ì': 'i', 'í': 'i', 'ỉ': 'i', 'ĩ': 'i', 'ị': 'i',
    'ò': 'o', 'ó': 'o', 'ỏ': 'o', 'õ': 'o', 'ọ': 'o',
    'ô': 'o', 'ố': 'o', 'ồ': 'o', 'ổ': 'o', 'ỗ': 'o', 'ộ': 'o',
    'ơ': 'o', 'ớ': 'o', 'ờ': 'o', 'ở': 'o', 'ỡ': 'o', 'ợ': 'o',
    'ù': 'u', 'ú': 'u', 'ủ': 'u', 'ũ': 'u', 'ụ': 'u',
    'ư': 'u', 'ứ': 'u', 'ừ': 'u', 'ử': 'u', 'ữ': 'u', 'ự': 'u',
    'ỳ': 'y', 'ý': 'y', 'ỷ': 'y', 'ỹ': 'y', 'ỵ': 'y'
}

def remove_vietnamese_accents(text):
    """
    Chuyển đổi toàn bộ chuỗi tiếng Việt có dấu thành không dấu (ASCII).
    Hỗ trợ cả Unicode dựng sẵn (NFC) và tổ hợp (NFD).
    """
    if not text:
        return ""
    # Chuẩn hóa về NFC trước
    s = unicodedata.normalize('NFC', text)
    result = []
    for ch in s:
        lower_ch = ch.lower()
        if lower_ch in VIETNAMESE_CHAR_MAP:
            mapped = VIETNAMESE_CHAR_MAP[lower_ch]
            result.append(mapped.upper() if ch.isupper() else mapped)
        else:
            result.append(ch)
    s = "".join(result)
    # Loại bỏ các dấu kết hợp còn sót lại (nếu có)
    s = unicodedata.normalize('NFD', s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != 'Mn')
    return s

def vietnamese_slugify(text, max_len=40, add_hash=True):
    """
    Sinh mã định danh chuẩn (Slug / MAT_ID) từ chuỗi tiếng Việt:
    - Loại bỏ dấu tiếng Việt.
    - Thay thế ký tự đặc biệt, khoảng trắng bằng dấu gạch dưới '_'.
    - Gộp các dấu gạch dưới liên tiếp thành một dấu '_'.
    - Thêm hash ngắn 4 ký tự nếu chuỗi bị cắt ngắn để chống xung đột (collision).
    """
    if not text:
        return "UNKNOWN"
        
    raw_clean = remove_vietnamese_accents(text.strip())
    # Thay thế các ký tự không phải chữ và số thành '_'
    slug = re.sub(r'[^a-zA-Z0-9]', '_', raw_clean)
    # Gộp nhiều dấu '_' liên tiếp thành 1
    slug = re.sub(r'_+', '_', slug).strip('_').upper()
    
    if not slug:
        slug = "ITEM"
        
    if len(slug) <= max_len:
        return slug
        
    if add_hash:
        # Băm MD5 4 ký tự từ chuỗi đã chuẩn hóa khoảng trắng để tránh xung đột
        normalized_text = re.sub(r'\s+', ' ', text.strip())
        hash_suffix = hashlib.md5(normalized_text.encode('utf-8')).hexdigest()[:4].upper()
        allowed_len = max_len - 5 # chừa chỗ cho '_' + 4 ký tự hash
        truncated = slug[:allowed_len].rstrip('_')
        return f"{truncated}_{hash_suffix}"
    else:
        return slug[:max_len].rstrip('_')

def detect_boq_header(ws):
    """
    Nhận diện dòng header và ánh xạ cột (col_map) của file BOQ khách hàng.
    Trả về (header_row, col_map, hidden_rows, hidden_cols, max_scan_col).
    Dùng chung cho cả SwindBOQParser và SwindExcelBuilder.
    """
    forced_q_key = getattr(ws, "_forced_qty_col", None)
    cached_hdr = getattr(ws, "_cached_boq_header", None)
    if cached_hdr is not None and cached_hdr[0] == forced_q_key:
        _, h_r, c_m, h_rows, h_cols, m_sc = cached_hdr
        return h_r, dict(c_m), set(h_rows), set(h_cols), m_sc

    hidden_cols = {col_let for col_let, dim in ws.column_dimensions.items() if dim.hidden}
    hidden_rows = {r_idx for r_idx, dim in ws.row_dimensions.items() if dim.hidden}
    max_scan_col = min(ws.max_column, 60)

    header_row = None
    col_map = {}

    aliases = get_subtype_rules().get("HEADER_ALIASES", {})
    NAME_KEYWORDS = tuple(aliases.get("name_keywords", ("sản phẩm", "hạng mục", "quy cách", "công tác", "công việc", "nội dung", "diễn giải", "mô tả", "tên vật tư", "danh mục", "tên hàng")))
    NAME_CAND_KEYWORDS = tuple(aliases.get("name_cand_keywords", ("sản phẩm", "hạng mục", "danh mục", "công tác", "công việc", "nội dung", "diễn giải", "mô tả", "tên vật tư", "tên hàng")))
    QTY_KEYWORDS = tuple(aliases.get("qty_keywords", ("số lượng", "q'ty", "q.ty", "qty", "khối lượng", "sl", "kl")))
    STT_KEYWORDS = tuple(aliases.get("stt_keywords", ("stt", "tt", "no", "no.")))
    SPEC_KEYWORDS = tuple(aliases.get("spec_keywords", ("thông số", "quy cách", "chi tiết")))
    MODEL_KEYWORDS = tuple(aliases.get("model_keywords", ("model", "mã hiệu", "ký hiệu", "mã sp")))
    BRAND_KEYWORDS = tuple(aliases.get("brand_keywords", ("thương hiệu", "nhãn hiệu", "xuất xứ")))
    SIZE_KEYWORDS = tuple(aliases.get("size_keywords", ("kích thước", "size")))
    THK_KEYWORDS = tuple(aliases.get("thickness_keywords", ("độ dày", "dày (mm)")))
    UNIT_KEYWORDS = tuple(aliases.get("unit_keywords", ("unit", "đơn vị", "đvt", "đv tính")))
    PRICE_KEYWORDS = tuple(aliases.get("price_keywords", ("đơn giá", "price")))
    TOTAL_KEYWORDS = tuple(aliases.get("total_keywords", ("thành tiền", "total")))
    REMARKS_KEYWORDS = tuple(aliases.get("remarks_keywords", ("ghi chú", "remarks")))

    forced_qty_col = getattr(ws, "_forced_qty_col", None)
    if forced_qty_col:
        hidden_cols.discard(get_column_letter(forced_qty_col))

    for r in range(1, min(25, ws.max_row + 1)):
        if r in hidden_rows:
            continue
        row_vals = {c: re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(ws.cell(r, c).value)).lower()).strip()
                    for c in range(1, max_scan_col + 1)
                    if get_column_letter(c) not in hidden_cols and ws.cell(r, c).value is not None}
        if len(row_vals) < 2:
            continue
        # Gom cả 2 dòng phụ ngay dưới (hỗ trợ header gộp nhiều dòng như phiếu ROP 01 / GIÓ-C3Z3 / ÔG MB Z3)
        block_vals = {c: [v] for c, v in row_vals.items()}
        for sub_r in range(r + 1, min(r + 3, ws.max_row + 1)):
            if sub_r in hidden_rows:
                continue
            for c_sub in range(1, max_scan_col + 1):
                if ws.cell(sub_r, c_sub).value is None:
                    continue
                v_s = re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(ws.cell(sub_r, c_sub).value)).lower()).strip()
                is_total_qty_col = any(pk in v_s for pk in ("tổng kl", "tổng khối lượng", "lũy kế đến hết kỳ này")) and "kỳ trước" not in v_s
                if get_column_letter(c_sub) in hidden_cols and not is_total_qty_col:
                    continue
                if v_s and not re.match(r'^\(\d+\)', v_s):
                    block_vals.setdefault(c_sub, []).append(v_s)

        first_price_col = min(
            (c for c, v in row_vals.items() if any(k in v for k in PRICE_KEYWORDS) or v in ("đg", "giá")),
            default=9999
        )

        def is_qty_label(v_txt, col_idx=0):
            if any(bad_t in v_txt for bad_t in ("ngày", "thời gian", "tiến độ", "đơn giá", "thành tiền")):
                return False
            if v_txt in ("tổng", "tổng số", "tổng cộng", "total qty", "total q'ty") and col_idx < first_price_col:
                return True
            return any(k in v_txt for k in QTY_KEYWORDS) or any(k in v_txt for k in ("cần về", "đặt hàng", "kl đặt", "sl đặt", "cung cấp đợt này", "yêu cầu cấp")) or bool(re.match(r'^lần\s*\d+$', v_txt))

        name_cols = [c for c, v in row_vals.items() if any(k in v for k in NAME_KEYWORDS)]
        qty_or_unit_cols = [
            c for c, vs in block_vals.items()
            if c not in name_cols and any(is_qty_label(v, c) or any(k in v for k in UNIT_KEYWORDS) for v in vs)
        ]
        has_qty_in_block = any(any(is_qty_label(v, c) for v in vs) for c, vs in block_vals.items() if c not in name_cols)
        if name_cols and qty_or_unit_cols and has_qty_in_block:
            header_row = r
            name_cands = []
            qty_cands = []
            qty_labels = {}
            last_hdr_r = r
            for scan_r in range(r, min(r + 3, ws.max_row + 1)):
                if scan_r in hidden_rows:
                    continue
                s_vals = {}
                for c in range(1, max_scan_col + 1):
                    if ws.cell(scan_r, c).value is None:
                        continue
                    v_tmp = re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(ws.cell(scan_r, c).value)).lower()).strip()
                    is_total_qty_col = any(pk in v_tmp for pk in ("tổng kl", "tổng khối lượng", "lũy kế đến hết kỳ này")) and "kỳ trước" not in v_tmp
                    if get_column_letter(c) in hidden_cols and not is_total_qty_col:
                        continue
                    s_vals[c] = v_tmp
                if scan_r > r and any(re.match(r'^(?:\(\d+\)|cột\s*\d+|column\s*\d+)$', s_vals.get(cc, "")) for cc in name_cols):
                    last_hdr_r = max(last_hdr_r, scan_r)
                    continue
                matched_sub = False
                for c, v in s_vals.items():
                    if scan_r == r and any(k in v for k in NAME_CAND_KEYWORDS):
                        if c not in name_cands: name_cands.append(c)
                    elif v in STT_KEYWORDS or v.startswith("stt"):
                        if "stt" not in col_map: col_map["stt"] = c
                    elif any(k in v for k in ("bích kết nối", "loại bích", "kiểu bích")):
                        col_map["flange"] = c
                        matched_sub = True
                    elif any(k in v for k in THK_KEYWORDS) or "độ dầy" in v or "độ dày tole" in v:
                        col_map["thickness"] = c
                        matched_sub = True
                    elif any(k in v for k in SPEC_KEYWORDS) or v in ("spec", "specs", "specification", "specifications"):
                        if "spec" not in col_map: col_map["spec"] = c
                        matched_sub = True
                    elif any(k in v for k in MODEL_KEYWORDS):
                        if "model" not in col_map: col_map["model"] = c
                    elif any(k in v for k in BRAND_KEYWORDS):
                        if "brand" not in col_map: col_map["brand"] = c
                    elif any(k in v for k in SIZE_KEYWORDS):
                        if "w" not in col_map: col_map["w"] = c
                    elif v in ("l", "chiều dài", "dài"):
                        if "l" not in col_map: col_map["l"] = c
                    elif any(k in v for k in UNIT_KEYWORDS):
                        if "unit" not in col_map: col_map["unit"] = c
                    elif is_qty_label(v, c):
                        if c not in name_cands:
                            if c not in qty_cands: qty_cands.append(c)
                            qty_labels[c] = f"{qty_labels[c]} {v}".strip() if c in qty_labels else v
                            matched_sub = True
                    elif c in qty_labels and scan_r > r:
                        qty_labels[c] = f"{qty_labels[c]} {v}".strip()
                    elif any(k in v for k in PRICE_KEYWORDS) or v in ("đg", "giá"):
                        if "price" not in col_map: col_map["price"] = c
                    elif any(k in v for k in TOTAL_KEYWORDS):
                        if "total" not in col_map: col_map["total"] = c
                    elif any(k in v for k in REMARKS_KEYWORDS):
                        if "remarks" not in col_map: col_map["remarks"] = c
                if scan_r > r and matched_sub:
                    last_hdr_r = max(last_hdr_r, scan_r)

            # Quét thêm các dòng phụ dưới header (để nhận diện cột con W, H, B1, H1, B2, H2, L, R(Ø), Độ lệch, dòng đánh số (1)(2))
            for sub_r in range(r + 1, min(r + 4, ws.max_row + 1)):
                if sub_r in hidden_rows:
                    continue
                row_has_col_idx = False
                row_has_sub_dim = False
                sub_has_w = False
                sub_has_h = False
                sub_has_w2 = False
                for c_sub in range(1, max_scan_col + 1):
                    if get_column_letter(c_sub) in hidden_cols:
                        continue
                    v_sub = unicodedata.normalize('NFC', str(ws.cell(sub_r, c_sub).value or "")).strip().lower()
                    if not v_sub:
                        continue
                    if c_sub in name_cols and re.match(r'^(?:\(\d+\)|cột\s*\d+|column\s*\d+)$', v_sub):
                        row_has_col_idx = True
                        continue
                    v_clean = re.sub(r'\s*\([^)]*\)', '', v_sub).strip()
                    if v_clean in ("w", "a1", "rộng"):
                        col_map["w"] = c_sub
                        sub_has_w = True
                        row_has_sub_dim = True
                    elif v_clean in ("w1", "b1"):
                        if not sub_has_w:
                            col_map["w"] = c_sub
                            sub_has_w = True
                        else:
                            col_map["w2"] = c_sub
                            sub_has_w2 = True
                        row_has_sub_dim = True
                    elif v_clean in ("h", "a2", "cao"):
                        col_map["h"] = c_sub
                        sub_has_h = True
                        row_has_sub_dim = True
                    elif v_clean == "h1":
                        if not sub_has_h:
                            col_map["h"] = c_sub
                            sub_has_h = True
                        else:
                            col_map["h2"] = c_sub
                        row_has_sub_dim = True
                    elif v_clean in ("w2", "w'"):
                        col_map["w2"] = c_sub
                        sub_has_w2 = True
                        row_has_sub_dim = True
                    elif v_clean == "b2":
                        if sub_has_w2:
                            col_map["h2"] = c_sub
                        else:
                            col_map["w2"] = c_sub
                            sub_has_w2 = True
                        row_has_sub_dim = True
                    elif v_clean in ("h2", "h'"):
                        col_map["h2"] = c_sub
                        row_has_sub_dim = True
                    elif v_clean in ("l", "l1", "dài"):
                        col_map["l"] = c_sub
                        row_has_sub_dim = True
                    elif v_clean in ("r", "r(ø)") or "r(ø)" in v_sub:
                        col_map["r_or_d"] = c_sub
                        row_has_sub_dim = True
                    elif "độ lệch" in v_sub:
                        col_map["offset"] = c_sub
                        row_has_sub_dim = True
                    elif "đơn giá" in v_sub and "price" not in col_map:
                        col_map["price"] = c_sub
                    elif "thành tiền" in v_sub and "total" not in col_map:
                        col_map["total"] = c_sub
                if row_has_col_idx or row_has_sub_dim:
                    last_hdr_r = max(last_hdr_r, sub_r)
            if col_map.get("spec") and col_map.get("spec") == col_map.get("w"):
                del col_map["spec"]

            header_row = last_hdr_r
            if name_cands:
                col_map["name"] = max(name_cands, key=lambda col_c: sum(1 for rr in range(header_row + 1, min(ws.max_row + 1, header_row + 25)) if ws.cell(rr, col_c).value not in (None, "")))
                for c_other in name_cands:
                    if c_other != col_map["name"]:
                        if "spec" not in col_map:
                            col_map["spec"] = c_other
                        elif "model" not in col_map and any(k in row_vals.get(c_other, "") for k in MODEL_KEYWORDS):
                            col_map["model"] = c_other
            if forced_qty_col:
                col_map["qty"] = forced_qty_col
            elif qty_cands:
                def _count_pos_qty(col_c):
                    cnt = 0
                    empty_s = 0
                    name_c = col_map.get("name", 2)
                    for rr in range(header_row + 1, ws.max_row + 1):
                        if rr in hidden_rows:
                            continue
                        v_q = ws.cell(rr, col_c).value
                        if isinstance(v_q, (int, float)) and v_q > 0:
                            cnt += 1
                            empty_s = 0
                        elif ws.cell(rr, name_c).value is None and v_q is None:
                            empty_s += 1
                            if empty_s >= 40:
                                break
                        else:
                            empty_s = 0
                    return cnt

                valid_qty = [
                    c for c in qty_cands
                    if not any(bad in qty_labels.get(c, "") for bad in ("tồn kho", "kỳ trước", "m2", "m²", "hợp đồng", "1 tầng", "một tầng", "điển hình"))
                ]
                if not valid_qty:
                    valid_qty = qty_cands
                order_qty = [
                    c for c in valid_qty
                    if any(ok in qty_labels.get(c, "") for ok in ("cần về", "đặt hàng", "kl đặt", "sl đặt", "cung cấp đợt này", "yêu cầu cấp"))
                    and _count_pos_qty(c) > 0
                ]
                total_qty = [
                    c for c in valid_qty
                    if any(tk in qty_labels.get(c, "") for tk in ("tổng kl", "tổng khối lượng", "tổng số lượng", "tổng sl", "lũy kế đến hết kỳ này"))
                    and _count_pos_qty(c) > 0
                ]
                std_qty = [
                    c for c in valid_qty
                    if not re.search(r'\blần\s*\d+\b', qty_labels.get(c, ""))
                    and any(pk in qty_labels.get(c, "") for pk in ("số lượng", "khối lượng", "q'ty", "q.ty", "qty", "sl", "kl"))
                    and _count_pos_qty(c) > 0
                ]
                lan_cands = [
                    (int(m.group(1)), c)
                    for c in valid_qty
                    for m in [re.search(r'\blần\s*(\d+)\b', qty_labels.get(c, ""))]
                    if m and _count_pos_qty(c) > 0
                ]
                if order_qty:
                    col_map["qty"] = max(order_qty, key=lambda col_c: (_count_pos_qty(col_c), col_c))
                    col_map["is_order_sheet"] = True
                elif lan_cands:
                    col_map["qty"] = max(lan_cands)[1]
                    col_map["is_order_sheet"] = True
                elif total_qty:
                    col_map["qty"] = max(total_qty, key=lambda col_c: (_count_pos_qty(col_c), col_c))
                    col_map["is_order_sheet"] = True
                elif std_qty:
                    col_map["qty"] = max(std_qty, key=lambda col_c: (_count_pos_qty(col_c), col_c))
                else:
                    col_map["qty"] = max(valid_qty, key=lambda col_c: (_count_pos_qty(col_c), col_c))
                if len(valid_qty) > 1 and not col_map.get("is_order_sheet"):
                    alt_cands = [c for c in valid_qty if c != col_map["qty"]]
                    if alt_cands:
                        col_map["qty_alt"] = alt_cands[0]
            break

    if header_row is None:
        unit_tokens = {"cái", "cai", "bộ", "bo", "chiếc", "chiec", "m", "m2", "m²", "md", "ống", "ong", "tấm", "tam", "cụm", "cum", "hộp", "hop", "lô", "lo", "kg", "set", "pcs", "ea", "nr", "lot", "nos", "unit", "units"}
        unit_col_counts = {}
        first_data_row = None
        for r in range(1, min(15, ws.max_row + 1)):
            if r in hidden_rows:
                continue
            for c in range(1, max_scan_col + 1):
                if get_column_letter(c) in hidden_cols:
                    continue
                val_s = str(ws.cell(r, c).value or "").strip().lower()
                if val_s in unit_tokens:
                    unit_col_counts[c] = unit_col_counts.get(c, 0) + 1
                    if first_data_row is None or r < first_data_row:
                        first_data_row = r
        if unit_col_counts and first_data_row is not None and first_data_row <= 10:
            u_col = max(unit_col_counts, key=unit_col_counts.get)
            name_col = 2
            best_name_score = -1
            for c in range(1, u_col):
                if get_column_letter(c) in hidden_cols:
                    continue
                score = sum(
                    1 for rr in range(first_data_row, min(ws.max_row + 1, first_data_row + 20))
                    if isinstance(ws.cell(rr, c).value, str) and len(ws.cell(rr, c).value.strip()) > 5
                )
                if score > best_name_score:
                    best_name_score = score
                    name_col = c
            qty_col = u_col + 1
            best_qty_score = -1
            for c in range(name_col + 1, max_scan_col + 1):
                if c == u_col or get_column_letter(c) in hidden_cols:
                    continue
                score = 0
                for rr in range(first_data_row, min(ws.max_row + 1, first_data_row + 20)):
                    v = ws.cell(rr, c).value
                    if isinstance(v, (int, float)) and v > 0:
                        score += 1
                    elif isinstance(v, str):
                        try:
                            if float(v.replace(",", "").strip()) > 0:
                                score += 1
                        except Exception:
                            pass
                if score > best_qty_score:
                    best_qty_score = score
                    qty_col = c
            header_row = first_data_row - 1
            while header_row > 1:
                u_v = ws.cell(header_row, u_col).value
                q_v = ws.cell(header_row, qty_col).value
                n_v = ws.cell(header_row, name_col).value
                n2_v = ws.cell(header_row, name_col + 1).value if name_col + 1 < u_col else None
                if u_v in (None, "") and q_v in (None, "") and (n_v not in (None, "") or n2_v not in (None, "")):
                    header_row -= 1
                else:
                    break
            col_map = {
                "stt": 1 if name_col > 1 else 0,
                "name": name_col,
                "unit": u_col,
                "qty": forced_qty_col if forced_qty_col else qty_col,
            }
            rem_col = max(u_col, qty_col) + 1
            if rem_col <= max_scan_col:
                col_map["remarks"] = rem_col
        else:
            for r in range(1, min(15, ws.max_row + 1)):
                if r in hidden_rows:
                    continue
                row_vals = {c: re.sub(r'\s+', ' ', str(ws.cell(r, c).value).lower()).strip()
                            for c in range(1, max_scan_col + 1)
                            if get_column_letter(c) not in hidden_cols and ws.cell(r, c).value is not None}
                n_cols = [c for c, v in row_vals.items() if any(k in v for k in NAME_CAND_KEYWORDS)]
                p_cols = [c for c, v in row_vals.items() if (any(k in v for k in PRICE_KEYWORDS) or v in ("giá", "gia")) and c not in n_cols]
                if n_cols and p_cols:
                    header_row = r
                    name_col = n_cols[0]
                    price_col = p_cols[0]
                    col_map = {
                        "stt": 1 if name_col > 1 else 0,
                        "name": name_col,
                        "price": price_col
                    }
                    break
            if header_row is None:
                dim_col_counts = {}
                first_dim_row = None
                for r in range(1, min(15, ws.max_row + 1)):
                    if r in hidden_rows:
                        continue
                    for c in range(1, max_scan_col + 1):
                        if get_column_letter(c) in hidden_cols:
                            continue
                        val_s = str(ws.cell(r, c).value or "").strip()
                        if re.search(r'(\d{2,4}\s*[xX*]\s*\d{2,4}|\b[dD]\d{2,4}\b)', val_s):
                            dim_col_counts[c] = dim_col_counts.get(c, 0) + 1
                            if first_dim_row is None or r < first_dim_row:
                                first_dim_row = r
                if dim_col_counts and first_dim_row is not None and first_dim_row <= 10:
                    name_col = max(dim_col_counts, key=dim_col_counts.get)
                    header_row = first_dim_row - 1
                    col_map = {
                        "stt": 1 if name_col > 1 else 0,
                        "name": name_col,
                    }
                    model_rx = re.compile(r'^(?:seg|sdg|eag|pag|sag|rag|fag|oag|mag|spg|wl|fal|oal|eal|sd|scd|sad|slb|ecg|vcd|obd|fd|mfd|mfsd|md|nrd|prd)(?:\s*\+\s*[a-z0-9]+)*$', re.IGNORECASE)
                    for c in range(name_col + 1, max_scan_col + 1):
                        if get_column_letter(c) in hidden_cols:
                            continue
                        col_vals = [
                            str(ws.cell(rr, c).value).strip()
                            for rr in range(first_dim_row, min(ws.max_row + 1, first_dim_row + 20))
                            if ws.cell(rr, c).value not in (None, "")
                        ]
                        if not col_vals:
                            continue
                        if "model" not in col_map and sum(1 for v in col_vals if model_rx.match(v)) >= max(1, len(col_vals) // 2):
                            col_map["model"] = c
                        elif "spec" not in col_map and sum(1 for v in col_vals if not v.replace(".", "").isdigit()) >= max(1, len(col_vals) // 2):
                            col_map["spec"] = c
                else:
                    # Fallback: bố cục chuẩn Swind (Row 11)
                    header_row = 11
                    col_map = {
                        "stt": 1, "name": 2, "desc": 3, "model": 4, "brand": 5,
                        "w": 6, "h": 7, "l": 11, "thickness": 12,
                        "unit": 13, "qty": 14, "price": 15, "total": 16, "remarks": 17
                    }

    ws._cached_boq_header = (forced_q_key, header_row, dict(col_map), set(hidden_rows), set(hidden_cols), max_scan_col)
    return header_row, col_map, hidden_rows, hidden_cols, max_scan_col


def extract_duct_fitting_geometry(text, default_w=0.0, default_h=0.0, default_d=0.0, default_w2=0.0, default_h2=0.0):
    """
    Giải mã chi tiết hình học của Ống gió & Phụ kiện ống gió từ chuỗi mô tả BOQ:
    Trả về dict:
      {
        "w1": float, "h1": float,
        "w2": float, "h2": float,
        "d1": float, "d2": float,
        "length": float,
        "radius": float,
        "offset_h": float,
        "angle": int,
        "thickness": float,
        "flange_type": str,
        "neck_count": int,
        "ei_rating": int,
        "is_insulated": bool,
        "material": str
      }
    """
    s = unicodedata.normalize('NFC', str(text or ""))
    s_lower = s.lower()

    w1, h1 = float(default_w or 0), float(default_h or 0)
    w2, h2 = float(default_w2 or 0), float(default_h2 or 0)
    d1, d2 = float(default_d or 0), 0.0
    length = 0.0
    has_explicit_l = False
    radius = 0.0
    offset_h = 0.0
    angle = 90
    thickness = 0.0
    neck_count = 1

    # 1. Tìm độ dày tôn (0.48, 0.58, 0.75, 0.95, 1.15, 1.5, hoặc 0.5, 0.6, 0.8, 1.0, 1.2)
    m_thk = re.search(
        r'(?:dày|day|độ dày|tôn|tole|tol|thk|δ)\s*[:=]?\s*(0[.,](?:47|48|5|50|57|58|6|60|64|7|75|77|8|80|95)|1[.,](?:0|00|15|2|20|48|5|50)|2[.,]0)\b'
        r'|\b(0[.,](?:48|58|75|95)|1[.,]15)\s*(?:mm)?\b',
        s_lower
    )
    if m_thk:
        thk_str = (m_thk.group(1) or m_thk.group(2)).replace(",", ".")
        thickness = float(thk_str)

    # 2. Tìm chiều dài L=... (hỗ trợ cả dạng xL250, xl200, L=500, L500, và lỗi gõ Ll180 -> L1180)
    m_len = re.search(r'(?:^|[\s_(*xX×,/+-])[lL]\s*[=:]?\s*(\d{2,4})\b', s)
    if m_len:
        length = float(m_len.group(1))
        has_explicit_l = True
    else:
        m_len_typo = re.search(r'(?:^|[\s_(*xX×,/+-])[lL][lI](\d{2,3})\b', s)
        if m_len_typo:
            length = float("1" + m_len_typo.group(1))
            has_explicit_l = True

    # 2b. Tìm độ lệch Z / E / H cho ống lượn Z / Z chuyển cao độ / Z chuyển hướng (hỗ trợ E320, /E80/, Z190, H300)
    m_off_e = re.search(r'(?:^|[\s_(*xX×,/+-])[eE]\s*[=:]?\s*(\d{2,4})\b', s)
    if m_off_e and float(m_off_e.group(1)) >= 40:
        offset_h = float(m_off_e.group(1))
    elif any(k in s_lower for k in ("chuyển cao độ", "chuyển hướng", "ống lượn", "lượn ống", "zét lượn", "zet luon", "zét")):
        m_off_h = re.search(r'(?:^|[\s_(*,/+-])[hH]\s*[=:]?\s*(\d{2,4})\b', s)
        if m_off_h and float(m_off_h.group(1)) >= 40:
            offset_h = float(m_off_h.group(1))
    if offset_h == 0:
        for m_z in re.finditer(r'(?:^|[\s_(*xX×,/+-])[zZ]\s*[=:]?\s*(\d{2,4})\b', s):
            if float(m_z.group(1)) >= 40:
                offset_h = float(m_z.group(1))
                break

    # Hỗ trợ dạng viết tắt không có chữ L sau dấu '/' như 'Z chuyển cao độ 100x300/E270/370' hoặc 'Ống gió bịt đầu 175x100/900'
    if length == 0:
        m_slash_l = re.search(r'/[eEzZ]\s*\d{2,4}\s*/\s*(\d{2,4})\b', s)
        if not m_slash_l:
            m_slash_l = re.search(r'\d{2,4}\s*[*xX×]\s*\d{2,4}\s*/\s*(\d{2,4})(?!\s*[*xX×\d])\b', s)
        if m_slash_l:
            length = float(m_slash_l.group(1))
            has_explicit_l = True

    # 3. Tìm bán kính R=... (hỗ trợ cả R125, R=200)
    m_rad = re.search(r'(?:^|[\s_(*xX×,/+-])[rR]\s*[=:]?\s*(\d{2,4})\b', s)
    if m_rad:
        radius = float(m_rad.group(1))

    # 4. Tìm góc cút / chếch (90, 60, 45, 30)
    m_ang = re.search(r'\b(90|60|45|30)\s*(?:độ|o|°|deg)?\b', s_lower)
    if m_ang:
        angle = int(m_ang.group(1))
    elif "chếch" in s_lower or "chech" in s_lower or "lơi" in s_lower:
        angle = 45

    # 5. Tìm đường kính tròn D / Ø / DN (hỗ trợ cả xD200, /D200, D300/200, D100)
    m_d_pair = re.search(
        r'(?:^|[\s_(*xX×,/+-])(?:[dD][nN]?|Ø|phi)\s*[=: ]?\s*(\d{2,4})\s*(?:[/\-]|xuống|ra|to)\s*(?:[dD][nN]?|Ø|phi)?\s*(\d{2,4})(?!\d)',
        s, re.IGNORECASE
    )
    d_matches = [
        float(m.group(1))
        for m in re.finditer(r'(?:^|[\s_(*xX×,/+-])(?:[dD][nN]?|Ø|phi)\s*[=: ]?\s*(\d{2,4})(?!\d)', s, re.IGNORECASE)
        if float(m.group(1)) >= 50
    ]
    if m_d_pair and float(m_d_pair.group(1)) >= 50 and float(m_d_pair.group(2)) >= 50:
        d1 = float(m_d_pair.group(1))
        d2 = float(m_d_pair.group(2))
    elif d_matches:
        d1 = d_matches[0]
        if len(d_matches) >= 2:
            d2 = d_matches[1]

    # Chuẩn bị chuỗi bóc tách kích thước vuông (loại bỏ xL..., xD..., R..., E..., Z... để không bắt nhầm vào WxH)
    rect_s = re.sub(r'(?:[*xX×,/+-]\s*)?[lLrReEzZ]\s*[=:]?\s*\d{2,4}\b', ' ', s)
    rect_s = re.sub(r'(?:[*xX×,/+-]\s*)?[lL][lI]\d{2,3}\b', ' ', rect_s)
    rect_s = re.sub(r'\b[vV]\s*\d{2,3}\s*[*xX×]\s*\d{2,3}(?:\s*[*xX×]\s*\d{1,2})?\b', ' ', rect_s)
    rect_s = re.sub(r'(?:[*xX×,/+-]\s*)?(?:[dD][nN]?|Ø|phi)\s*[=:]?\s*\d{2,4}(?:\s*[/\-]\s*(?:[dD][nN]?|Ø|phi)?\s*\d{2,4})?\b', ' ', rect_s, flags=re.IGNORECASE)

    # 6. Tìm kích thước 4 chiều (W1xH1 / W2xH2 hoặc w700xH300xW550xH300) hoặc 3 chiều (W1xH1xL) hoặc 2 chiều (W1xH1)
    m_4dim = re.search(
        r'[wWaA]?\s*(\d{2,4})\s*[*xX×]\s*[hHbBwW]?\s*(\d{2,4})\s*(?:[/\-]|[*xX×]|xuống|sang|ra|to)\s*[wWaA]?\s*(\d{2,4})\s*[*xX×]\s*[hHbBwW]?\s*(\d{2,4})',
        rect_s
    )
    if m_4dim:
        w1, h1 = float(m_4dim.group(1)), float(m_4dim.group(2))
        w2, h2 = float(m_4dim.group(3)), float(m_4dim.group(4))
    else:
        m_3dim = re.search(r'[wWaA]?\s*(\d{2,4})\s*[*xX×]\s*[hHbBwW]?\s*(\d{2,4})\s*[*xX×]\s*[lL]?\s*(\d{2,4})', rect_s)
        if m_3dim:
            v1, v2, v3 = float(m_3dim.group(1)), float(m_3dim.group(2)), float(m_3dim.group(3))
            if v1 >= 50 and v2 >= 50:
                if w1 == 0 and h1 == 0:
                    w1, h1 = v1, v2
                if length == 0 and v3 >= 50:
                    length = v3
                    has_explicit_l = True
        else:
            pairs = []
            for m in re.finditer(r'[wWaA]?\s*(\d{2,4})\s*(?:mm)?\s*[*xX×]\s*[hHbBwW]?\s*(\d{2,4})', rect_s):
                v1, v2 = float(m.group(1)), float(m.group(2))
                if v1 >= 50 and v2 >= 50:
                    pairs.append((v1, v2))
                elif v1 >= 100 and m.group(2) == "00":
                    pairs.append((v1, 100.0 if v1 <= 450 else v1))
            if pairs:
                if w1 == 0 and h1 == 0:
                    w1, h1 = pairs[0]
                if len(pairs) >= 2 and w2 == 0 and h2 == 0:
                    w2, h2 = pairs[1]

    m_neck_cnt = re.search(r'(?<![a-zA-Z0-9Øø])(\d{1,2})\s*(?:cổ|lỗ)', s_lower)
    if m_neck_cnt:
        neck_count = int(m_neck_cnt.group(1))

    flange_type = "TDC"
    if any(k in s_lower for k in ("nẹp c", "nep c", "bích c", "bich c")) or re.search(r'\bbích\s+c\b', s_lower):
        flange_type = "NEP_C"
    elif any(k in s_lower for k in ("tdc", "bích tdc", "bich tdc", "tdf")):
        flange_type = "TDC"

    ei_rating = 0
    m_ei = re.search(r'\bei\s*[-:]?\s*(\d{2,3})\b', s_lower)
    if m_ei:
        ei_rating = int(m_ei.group(1))

    is_insulated = any(k in s_lower for k in ("có bảo ôn", "bọc bảo ôn", "bọc pe", "cách nhiệt")) and not any(k in s_lower for k in ("không bảo ôn", "ko bảo ôn", "không cách"))

    material = "TON_HOA_SEN"
    if "inox" in s_lower or "sus" in s_lower:
        material = "INOX_SUS304"
    elif "phương nam" in s_lower:
        material = "TON_PHUONG_NAM"

    return {
        "w1": w1, "h1": h1,
        "w2": w2, "h2": h2,
        "d1": d1, "d2": d2,
        "length": length,
        "has_explicit_l": has_explicit_l,
        "radius": radius,
        "offset_h": offset_h,
        "angle": angle,
        "thickness": thickness,
        "flange_type": flange_type,
        "neck_count": neck_count,
        "ei_rating": ei_rating,
        "is_insulated": is_insulated,
        "material": material
    }

_CACHED_PRICE_MAP = None
_CACHED_BOM_RULES = None
_CACHED_SUBTYPE_RULES = None

def get_subtype_rules(force_reload=False):
    """
    Đọc bộ quy tắc phân loại sản phẩm, từ khóa nhận diện & quy tắc Combo từ data/master_data/subtype_rules.json.
    """
    global _CACHED_SUBTYPE_RULES
    if _CACHED_SUBTYPE_RULES is not None and not force_reload:
        return _CACHED_SUBTYPE_RULES

    import os
    import json
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    rules_path = os.path.join(base_dir, "data", "master_data", "subtype_rules.json")
    if os.path.exists(rules_path):
        with open(rules_path, "r", encoding="utf-8") as f:
            _CACHED_SUBTYPE_RULES = json.load(f)
    else:
        _CACHED_SUBTYPE_RULES = {}
    return _CACHED_SUBTYPE_RULES

def get_master_price_map(force_reload=False):
    """
    Đọc trực tiếp bảng đơn giá vật tư từ bộ file Master Data đã có sẵn (data/master_data/).
    Luôn ưu tiên phiên bản giá mới nhất trong data/master_data/prices/ (nếu có update giá).
    """
    global _CACHED_PRICE_MAP
    if _CACHED_PRICE_MAP is not None and not force_reload:
        return _CACHED_PRICE_MAP

    import os
    import glob
    import json
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    master_dir = os.path.join(base_dir, "data", "master_data")
    prices_dir = os.path.join(master_dir, "prices")

    db_files = sorted(glob.glob(os.path.join(prices_dir, "master_database_price_*.json")))
    db_path = db_files[-1] if db_files else os.path.join(master_dir, "master_database_price.json")
    price_map = {}
    if os.path.exists(db_path):
        with open(db_path, "r", encoding="utf-8") as f:
            for item in json.load(f):
                price_map[item["MAT_ID"]] = float(item.get("Price") or 0.0)

    com_files = sorted(glob.glob(os.path.join(prices_dir, "master_commercial_price_*.json")))
    com_path = com_files[-1] if com_files else os.path.join(master_dir, "master_commercial_price.json")
    if os.path.exists(com_path):
        with open(com_path, "r", encoding="utf-8") as f:
            for item in json.load(f):
                price_map[item["MAT_ID"]] = float(item.get("Price") or 0.0)

    _CACHED_PRICE_MAP = price_map
    return _CACHED_PRICE_MAP

def get_master_bom_rules(force_reload=False):
    """
    Đọc trực tiếp bộ quy tắc & hệ số định mức chuẩn từ data/master_data/master_bom_rules.json (đã giải mã 1 lần).
    """
    global _CACHED_BOM_RULES
    if _CACHED_BOM_RULES is not None and not force_reload:
        return _CACHED_BOM_RULES

    import os
    import json
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    rules_path = os.path.join(base_dir, "data", "master_data", "master_bom_rules.json")
    if os.path.exists(rules_path):
        with open(rules_path, "r", encoding="utf-8") as f:
            _CACHED_BOM_RULES = json.load(f)
    else:
        _CACHED_BOM_RULES = {}
    return _CACHED_BOM_RULES

if __name__ == "__main__":
    test_cases = [
        "Lưới bảo vệ",
        "Quạt hướng trục CEF-B2-Z1-01~02 5290/9430 l/s @ 500/1350 Pa, chống cháy 300°C trong 2h",
        "Cửa gió nan Z 600x600",
        "Chi phí cắt đục, đóng lưới trám trét tường",
        "Lớp chống cháy EI30",
        "Vật tư phụ, phụ kiện"
    ]
    for tc in test_cases:
        print(f"'{tc}'\n -> Slug: '{vietnamese_slugify(tc)}'\n")
