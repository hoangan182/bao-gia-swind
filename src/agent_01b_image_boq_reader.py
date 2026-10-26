# -*- coding: utf-8 -*-
"""
SWIND-AGENT-01B: IMAGE BOQ READER (OCR & 2D SPATIAL TABLE RECONSTRUCTOR)
Tự động đọc các bức ảnh chụp bảng khối lượng (BOQ) khách gửi trong thư mục 'input picture/',
khôi phục cấu trúc bảng (STT, Danh mục sản phẩm, Đơn vị, Số lượng, Ghi chú) và xuất ra
file Excel chuẩn trong thư mục 'input/' để chạy chung luồng với BOQ Excel truyền thống.
"""

import os
import re
import sys
import json
import zipfile
import tempfile
import subprocess
from datetime import datetime
from xml.sax.saxutils import escape as xml_escape


SWIFT_VISION_OCR_SCRIPT = r"""
import Foundation
import Vision
import AppKit

let args = CommandLine.arguments
guard args.count > 1 else { exit(1) }

var allResults: [[String: Any]] = []

for i in 1..<args.count {
    let path = args[i]
    guard let image = NSImage(contentsOfFile: path),
          let cgImage = image.cgImage(forProposedRect: nil, context: nil, hints: nil) else {
        continue
    }

    let request = VNRecognizeTextRequest()
    request.recognitionLevel = .accurate
    request.recognitionLanguages = ["vi-VN", "en-US"]
    request.usesLanguageCorrection = true

    let handler = VNImageRequestHandler(cgImage: cgImage, options: [:])
    try? handler.perform([request])

    var items: [[String: Any]] = []
    if let observations = request.results {
        for obs in observations {
            if let top = obs.topCandidates(1).first {
                let box = obs.boundingBox
                items.append([
                    "text": top.string,
                    "confidence": top.confidence,
                    "x": box.origin.x,
                    "y": box.origin.y,
                    "w": box.size.width,
                    "h": box.size.height
                ])
            }
        }
    }
    allResults.append([
        "file": (path as NSString).lastPathComponent,
        "path": path,
        "width": cgImage.width,
        "height": cgImage.height,
        "items": items
    ])
}

if let data = try? JSONSerialization.data(withJSONObject: allResults, options: [.prettyPrinted]),
   let str = String(data: data, encoding: .utf8) {
    print(str)
}
"""


class Agent01B_ImageBOQReader:
    """
    Agent 01B: Đọc BOQ từ ảnh chụp màn hình khách gửi (Zalo, Viber, Email...)
    và chuyển đổi thành bảng Excel BOQ chuẩn Swind.
    """

    VALID_IMAGE_EXTS = ('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tiff')

    UNIT_MAP = {
        'cái': 'Cái',
        'cai': 'Cái',
        'bộ': 'Bộ',
        'bo': 'Bộ',
        'chiếc': 'Cái',
        'chiec': 'Cái',
        'cụm': 'Cụm',
        'cum': 'Cụm',
        'tấm': 'Tấm',
        'tam': 'Tấm',
        'm': 'm',
        'md': 'md',
        'm2': 'm2',
        'm²': 'm2',
        'kg': 'kg',
        'hộp': 'Hộp',
        'ống': 'Ống',
        'lô': 'Lô',
    }

    UI_NOISE_PATTERNS = [
        r'^(insert|page layout|formulas|data|review|view|developer|help|home|file)$',
        r'^(ready|accessibility|good to go|sheet\d*|danh mục vật tư)$',
        r'^(fx|b\d{1,3}|a\d{1,3}|c\d{1,3})$',
        r'^[a-z]$',
        r'^(stt|đvt|đơn vị|số lượng|khối lượng|ghi chú|đơn giá|thành tiền|mô tả|quy cách|tên vật tư|ợng)$',
    ]

    OCR_HVAC_CORRECTIONS = [
        (r'\bLàp\s+đặt\b', 'Lắp đặt'),
        (r'\blàp\s+đặt\b', 'lắp đặt'),
        (r'\bCửra\b', 'Cửa'),
        (r'\bcửra\b', 'cửa'),
        (r'\bCưa\b', 'Cửa'),
        (r'\bcưa\b', 'cửa'),
        (r'\bnan\s+thăng\b', 'nan thẳng'),
        (r'\bkt\s+cô\b', 'kt cổ'),
        (r'\bKT\s+cô\b', 'KT cổ'),
        (r'\bđiều\s+chinh\b', 'điều chỉnh'),
        (r'\blưu\s+lương\b', 'lưu lượng'),
        (r'\bchăn\s+côn\s+trùng\b', 'chắn côn trùng'),
        (r'\blưới\s+chăn\b', 'lưới chắn'),
        (r'\bcôn\s+trủng\b', 'côn trùng'),
        (r'\bcôn\s+trung\b', 'côn trùng'),
    ]

    def __init__(self, workspace_dir):
        self.workspace_dir = workspace_dir
        self.default_image_dir = os.path.join(workspace_dir, "input picture")
        self.default_input_dir = os.path.join(workspace_dir, "input")

    def list_customer_images(self, image_dir=None):
        target_dir = image_dir or self.default_image_dir
        if not os.path.exists(target_dir):
            return []
        files = []
        for f in sorted(os.listdir(target_dir)):
            if f.startswith('.'):
                continue
            if f.lower().endswith(self.VALID_IMAGE_EXTS):
                files.append(os.path.join(target_dir, f))
        return files

    def run_vision_ocr(self, image_paths):
        """Chạy macOS Vision OCR (VNRecognizeTextRequest) cho danh sách ảnh."""
        if not image_paths:
            return []

        cache_dir = "/tmp/swift_module_cache"
        os.makedirs(cache_dir, exist_ok=True)

        with tempfile.NamedTemporaryFile(mode="w", suffix=".swift", dir="/tmp", delete=False, encoding="utf-8") as sf:
            sf.write(SWIFT_VISION_OCR_SCRIPT)
            swift_path = sf.name

        try:
            cmd = ["/usr/bin/swift", "-module-cache-path", cache_dir, swift_path] + list(image_paths)
            proc = subprocess.run(cmd, capture_output=True, text=True, check=True)
            raw_out = proc.stdout.strip()
            if not raw_out:
                return []
            return json.loads(raw_out)
        finally:
            if os.path.exists(swift_path):
                try:
                    os.remove(swift_path)
                except OSError:
                    pass

    def _is_ui_noise(self, text, y_center):
        t = text.strip()
        if not t:
            return True
        t_low = t.lower()
        for pat in self.UI_NOISE_PATTERNS:
            if re.match(pat, t_low):
                return True
        if t.startswith('+') and len(t) <= 3:
            return True
        return False

    def _normalize_unit(self, text):
        t_clean = text.strip().lower().strip('.,:;|-')
        return self.UNIT_MAP.get(t_clean)

    def _clean_hvac_ocr_text(self, text):
        if not text:
            return ""
        res = re.sub(r'\s+', ' ', text).strip()
        for pat, repl in self.OCR_HVAC_CORRECTIONS:
            res = re.sub(pat, repl, res)
        return res

    def _parse_numeric_qty(self, text):
        t = text.strip().replace(' ', '')
        if not re.match(r'^\d+(?:[.,]\d+)?$', t):
            return None
        try:
            val = float(t.replace(',', '.'))
            if val <= 0 or val > 100000:
                return None
            return val
        except ValueError:
            return None

    def reconstruct_rows_from_image(self, image_ocr_data):
        """
        Thuật toán 2D Spatial Table Reconstruction:
        1. Tìm các cặp neo (Unit token, Quantity token) trên cùng 1 hàng ngang.
        2. Phân vùng không gian theo trục Y (Voronoi vertical band) cho từng dòng BOQ.
        3. Gom các dòng mô tả (kể cả mô tả bị ngắt thành 2-3 dòng) bên trái cột ĐVT
           và gom cột Ghi chú (nếu có) bên phải cột Số lượng.
        """
        file_name = image_ocr_data.get("file", "")
        raw_items = image_ocr_data.get("items", [])
        if not raw_items:
            return []

        tokens = []
        for idx, it in enumerate(raw_items):
            txt = str(it.get("text", "")).strip()
            if not txt:
                continue
            x = float(it.get("x", 0.0))
            y = float(it.get("y", 0.0))
            w = float(it.get("w", 0.0))
            h = float(it.get("h", 0.0))
            y_center = y + (h / 2.0)
            if self._is_ui_noise(txt, y_center):
                continue
            tokens.append({
                "id": idx,
                "text": txt,
                "x": x,
                "y": y,
                "w": w,
                "h": h,
                "x_right": x + w,
                "y_center": y_center,
                "conf": float(it.get("confidence", 1.0)),
            })

        unit_candidates = []
        for tok in tokens:
            norm_u = self._normalize_unit(tok["text"])
            if norm_u and tok["x"] > 0.25:
                unit_candidates.append((tok, norm_u))

        row_anchors = []
        used_token_ids = set()
        for u_tok, norm_u in unit_candidates:
            best_qty_tok = None
            best_qty_val = None
            best_dist_x = 999.0
            for q_tok in tokens:
                if q_tok["id"] == u_tok["id"] or q_tok["id"] in used_token_ids:
                    continue
                if q_tok["x"] <= u_tok["x"]:
                    continue
                if abs(q_tok["y_center"] - u_tok["y_center"]) > 0.06:
                    continue
                q_val = self._parse_numeric_qty(q_tok["text"])
                if q_val is not None:
                    dist_x = q_tok["x"] - u_tok["x"]
                    if dist_x < best_dist_x:
                        best_dist_x = dist_x
                        best_qty_tok = q_tok
                        best_qty_val = q_val

            if best_qty_tok is not None:
                used_token_ids.add(u_tok["id"])
                used_token_ids.add(best_qty_tok["id"])
                row_anchors.append({
                    "unit_tok": u_tok,
                    "unit_str": norm_u,
                    "qty_tok": best_qty_tok,
                    "qty_val": best_qty_val,
                    "y_anchor": (u_tok["y_center"] + best_qty_tok["y_center"]) / 2.0,
                    "x_unit": u_tok["x"],
                    "x_qty_right": best_qty_tok["x_right"],
                })

        if not row_anchors:
            sorted_toks = sorted(tokens, key=lambda t: (-t["y_center"], t["x"]))
            text_lines = []
            for tok in sorted_toks:
                if not text_lines or abs(tok["y_center"] - text_lines[-1]["y_center"]) > 0.025:
                    text_lines.append({"y_center": tok["y_center"], "tokens": [tok]})
                else:
                    text_lines[-1]["tokens"].append(tok)

            merged_descriptions = []
            start_kw_re = re.compile(
                r'^(?:\d+[\.\)\s]+)?(?:lắp\s+đặt|làp\s+đặt|cung\s+cấp|cửa|cưa|van|ống|tiêu\s+âm|côn\s+(?:thu|giảm|vuông|tròn|lệch|đồng)|cút|tê|chạc|hộp|măng|chân\s+rẽ|bịt|mặt\s+nạ|khuếch\s+tán|slot|louver|vcd|fd|mfd|md|obd|prd|nrd)\b',
                re.IGNORECASE
            )
            dim_re = re.compile(r'\b\d{2,4}\s*[xX*×]\s*\d{2,4}(?:mm|\b)|[ØD]\s*=?\s*\d{2,4}(?:mm|\b)', re.IGNORECASE)

            for ln in text_lines:
                ln["tokens"].sort(key=lambda t: t["x"])
                line_str = self._clean_hvac_ocr_text(" ".join(t["text"].strip() for t in ln["tokens"]))
                if not line_str:
                    continue
                is_new_item = bool(start_kw_re.search(line_str) or dim_re.search(line_str))
                if merged_descriptions and not is_new_item:
                    merged_descriptions[-1] = self._clean_hvac_ocr_text(merged_descriptions[-1] + " " + line_str)
                else:
                    merged_descriptions.append(line_str)

            return [
                {
                    "source_image": file_name,
                    "stt_goc": str(i),
                    "description": desc,
                    "unit": "Cái",
                    "qty": 1,
                    "note": "",
                }
                for i, desc in enumerate(merged_descriptions, start=1)
            ]

        row_anchors.sort(key=lambda r: r["y_anchor"], reverse=True)

        n_rows = len(row_anchors)
        if n_rows > 1:
            gaps = [row_anchors[i]["y_anchor"] - row_anchors[i + 1]["y_anchor"] for i in range(n_rows - 1)]
            avg_half_gap = max(0.08, (sum(gaps) / len(gaps)) * 0.65)
        else:
            avg_half_gap = 0.20

        extracted_rows = []
        for i, anchor in enumerate(row_anchors):
            if i == 0:
                y_top = min(0.54 if any(t["text"].lower() == "insert" for t in raw_items) else 1.0,
                            anchor["y_anchor"] + avg_half_gap)
            else:
                y_top = (row_anchors[i - 1]["y_anchor"] + anchor["y_anchor"]) / 2.0

            if i == n_rows - 1:
                y_bot = max(0.0, anchor["y_anchor"] - avg_half_gap)
            else:
                y_bot = (anchor["y_anchor"] + row_anchors[i + 1]["y_anchor"]) / 2.0

            band_tokens = [
                t for t in tokens
                if t["id"] not in used_token_ids and (y_bot <= t["y_center"] <= y_top)
            ]

            left_tokens = [t for t in band_tokens if t["x"] < anchor["x_unit"] - 0.02]
            right_tokens = [t for t in band_tokens if t["x"] > anchor["x_qty_right"] + 0.005]

            left_tokens.sort(key=lambda t: (-round(t["y_center"], 2), t["x"]))

            stt_candidates = []
            desc_tokens = []
            for t in left_tokens:
                if re.match(r'^\d{1,3}$', t["text"].strip()) and t["x"] < 0.065:
                    stt_candidates.append(t)
                else:
                    desc_tokens.append(t)

            stt_goc = ""
            if stt_candidates:
                stt_candidates.sort(key=lambda t: t["x"])
                chosen_stt = stt_candidates[-1]
                if chosen_stt["x"] >= 0.025:
                    stt_goc = chosen_stt["text"].strip()

            desc_tokens.sort(key=lambda t: (-t["y_center"], t["x"]))
            description = self._clean_hvac_ocr_text(" ".join(t["text"].strip() for t in desc_tokens))

            right_tokens.sort(key=lambda t: (-t["y_center"], t["x"]))
            note_parts = []
            for t in right_tokens:
                if re.match(r'^[\d.,\s-]+$', t["text"].strip()):
                    continue
                note_parts.append(t["text"].strip())
            note_str = self._clean_hvac_ocr_text(" ".join(note_parts))

            if description:
                extracted_rows.append({
                    "source_image": file_name,
                    "stt_goc": stt_goc,
                    "description": description,
                    "unit": anchor["unit_str"],
                    "qty": anchor["qty_val"],
                    "note": note_str,
                })

        return extracted_rows

    def write_boq_excel(self, rows, output_xlsx_path):
        """
        Xuất danh sách dòng BOQ từ ảnh ra file Excel chuẩn (.xlsx) tại `output_xlsx_path`.
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_xlsx_path)), exist_ok=True)
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "BOQ_TU_ANH"

            title_font = Font(name="Times New Roman", size=13, bold=True, color="1F4E79")
            header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            cell_font = Font(name="Times New Roman", size=11)
            note_font = Font(name="Times New Roman", size=11, bold=True, color="C00000")
            note_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_border = Border(
                left=Side(style="thin", color="BFBFBF"),
                right=Side(style="thin", color="BFBFBF"),
                top=Side(style="thin", color="BFBFBF"),
                bottom=Side(style="thin", color="BFBFBF"),
            )

            ws["A1"] = "BẢNG TỔNG HỢP KHỐI LƯỢNG BÓC TÁCH TỪ ẢNH KHÁCH HÀNG (SWIND OCR)"
            ws["A1"].font = title_font
            ws.merge_cells("A1:F1")

            headers = ["STT", "Danh mục vật tư / Quy cách kỹ thuật", "ĐVT", "Khối lượng", "Ghi chú", "Nguồn ảnh"]
            ws.append([])
            ws.append(headers)
            header_row_idx = 3
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=header_row_idx, column=col_idx)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = thin_border

            for idx, r in enumerate(rows, start=1):
                row_num = header_row_idx + idx
                qty_val = r["qty"]
                if isinstance(qty_val, float) and qty_val.is_integer():
                    qty_val = int(qty_val)
                vals = [
                    idx,
                    r["description"],
                    r["unit"],
                    qty_val,
                    r.get("note", ""),
                    r.get("source_image", ""),
                ]
                for col_idx, val in enumerate(vals, start=1):
                    c = ws.cell(row=row_num, column=col_idx, value=val)
                    c.font = cell_font
                    c.border = thin_border
                    if col_idx in (1, 3, 4):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    if col_idx == 5 and val:
                        c.font = note_font
                        c.fill = note_fill

            ws.column_dimensions["A"].width = 8
            ws.column_dimensions["B"].width = 65
            ws.column_dimensions["C"].width = 10
            ws.column_dimensions["D"].width = 14
            ws.column_dimensions["E"].width = 42
            ws.column_dimensions["F"].width = 28

            wb.save(output_xlsx_path)
            return output_xlsx_path
        except ImportError:
            return self._write_openxml_xlsx_fallback(rows, output_xlsx_path)

    def _write_openxml_xlsx_fallback(self, rows, output_xlsx_path):
        """Bộ tạo file .xlsx chuẩn OpenXML không cần thư viện ngoài."""
        shared_strings = []
        str_map = {}

        def get_si(s):
            s = str(s)
            if s not in str_map:
                str_map[s] = len(shared_strings)
                shared_strings.append(s)
            return str_map[s]

        sheet_rows_xml = []
        title_si = get_si("BẢNG TỔNG HỢP KHỐI LƯỢNG BÓC TÁCH TỪ ẢNH KHÁCH HÀNG (SWIND OCR)")
        sheet_rows_xml.append(f'<row r="1"><c r="A1" t="s"><v>{title_si}</v></c></row>')

        headers = ["STT", "Danh mục vật tư / Quy cách kỹ thuật", "ĐVT", "Khối lượng", "Ghi chú", "Nguồn ảnh"]
        col_letters = ["A", "B", "C", "D", "E", "F"]
        h_cells = []
        for c_let, h_text in zip(col_letters, headers):
            h_si = get_si(h_text)
            h_cells.append(f'<c r="{c_let}3" t="s"><v>{h_si}</v></c>')
        sheet_rows_xml.append(f'<row r="3">{"".join(h_cells)}</row>')

        for idx, r in enumerate(rows, start=1):
            r_num = idx + 3
            qty_val = r["qty"]
            if isinstance(qty_val, float) and qty_val.is_integer():
                qty_val = int(qty_val)
            c_stt = f'<c r="A{r_num}" t="n"><v>{idx}</v></c>'
            c_desc = f'<c r="B{r_num}" t="s"><v>{get_si(r["description"])}</v></c>'
            c_unit = f'<c r="C{r_num}" t="s"><v>{get_si(r["unit"])}</v></c>'
            c_qty = f'<c r="D{r_num}" t="n"><v>{qty_val}</v></c>'
            c_note = f'<c r="E{r_num}" t="s"><v>{get_si(r.get("note", ""))}</v></c>'
            c_src = f'<c r="F{r_num}" t="s"><v>{get_si(r.get("source_image", ""))}</v></c>'
            sheet_rows_xml.append(f'<row r="{r_num}">{c_stt}{c_desc}{c_unit}{c_qty}{c_note}{c_src}</row>')

        sst_items = "".join(f"<si><t>{xml_escape(s)}</t></si>" for s in shared_strings)
        sst_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(shared_strings)}" uniqueCount="{len(shared_strings)}">'
            f'{sst_items}</sst>'
        )

        sheet_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetData>{"".join(sheet_rows_xml)}</sheetData>'
            '</worksheet>'
        )

        workbook_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="BOQ_TU_ANH" sheetId="1" r:id="rId1"/></sheets>'
            '</workbook>'
        )

        wb_rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            '</Relationships>'
        )

        styles_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="1"><font><sz val="11"/><name val="Times New Roman"/></font></fonts>'
            '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>'
            '</styleSheet>'
        )

        root_rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>'
        )

        content_types_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
            '</Types>'
        )

        with zipfile.ZipFile(output_xlsx_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types_xml)
            zf.writestr("_rels/.rels", root_rels_xml)
            zf.writestr("xl/workbook.xml", workbook_xml)
            zf.writestr("xl/_rels/workbook.xml.rels", wb_rels_xml)
            zf.writestr("xl/styles.xml", styles_xml)
            zf.writestr("xl/sharedStrings.xml", sst_xml)
            zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)

        return output_xlsx_path

    def process_image_folder(self, image_dir=None, output_xlsx_path=None, output_json_path=None, split_per_image=True):
        """
        Đọc toàn bộ ảnh trong `image_dir` (mặc định là `input picture/`),
        trích xuất các dòng BOQ và lưu thành file Excel chuẩn trong `input/`.
        """
        if image_dir and os.path.isfile(image_dir) and image_dir.lower().endswith(self.VALID_IMAGE_EXTS):
            img_paths = [os.path.abspath(image_dir)]
        else:
            img_paths = self.list_customer_images(image_dir)
        if not img_paths:
            return {"status": "NO_IMAGES", "rows": [], "excel_path": None, "excel_paths": []}

        print(f"📸 [Agent 01B] Đang đọc OCR {len(img_paths)} bức ảnh khách gửi trong '{image_dir or self.default_image_dir}'...")
        ocr_results = self.run_vision_ocr(img_paths)

        all_rows = []
        excel_paths = []
        per_image_groups = []

        for idx, img_data in enumerate(ocr_results, start=1):
            img_rows = self.reconstruct_rows_from_image(img_data)
            print(f"   + Ảnh {idx:02d} ('{img_data.get('file')}'): Trích xuất được {len(img_rows)} dòng vật tư.")
            all_rows.extend(img_rows)
            if img_rows:
                per_image_groups.append((idx, img_data.get("file", f"img_{idx}"), img_rows))

        if not all_rows:
            return {"status": "EMPTY_OCR", "rows": [], "excel_path": None, "excel_paths": []}

        if split_per_image:
            old_merged = os.path.join(self.default_input_dir, "BOQ_TU_ANH_KHACH_GUI.xlsx")
            if os.path.exists(old_merged):
                try:
                    os.remove(old_merged)
                except OSError:
                    pass

            for idx, img_file, img_rows in per_image_groups:
                stem = os.path.splitext(os.path.basename(img_file))[0]
                if image_dir and os.path.isfile(image_dir) and not re.match(r'^z\d{8,}', stem, re.IGNORECASE):
                    out_path = os.path.join(self.default_input_dir, f"{stem}.xlsx")
                else:
                    out_path = os.path.join(self.default_input_dir, f"BOQ_ANH_KHACH_{idx:02d}.xlsx")
                self.write_boq_excel(img_rows, out_path)
                excel_paths.append(out_path)
                print(f"✅ [Agent 01B] Khách {idx:02d} ({img_file}) -> Xuất {len(img_rows)} dòng ra: {out_path}")
        else:
            if not output_xlsx_path:
                output_xlsx_path = os.path.join(self.default_input_dir, "BOQ_TU_ANH_KHACH_GUI.xlsx")
            self.write_boq_excel(all_rows, output_xlsx_path)
            excel_paths.append(output_xlsx_path)
            print(f"✅ [Agent 01B] Đã kết xuất {len(all_rows)} dòng BOQ từ ảnh ra file Excel chuẩn: {output_xlsx_path}")

        if output_json_path:
            os.makedirs(os.path.dirname(os.path.abspath(output_json_path)), exist_ok=True)
            with open(output_json_path, "w", encoding="utf-8") as f:
                json.dump({
                    "timestamp": datetime.now().isoformat(),
                    "total_images": len(img_paths),
                    "total_rows": len(all_rows),
                    "output_excels": excel_paths,
                    "rows": all_rows,
                }, f, ensure_ascii=False, indent=2)

        return {
            "status": "SUCCESS",
            "rows": all_rows,
            "excel_path": excel_paths[0] if excel_paths else None,
            "excel_paths": excel_paths,
            "image_count": len(img_paths),
        }


if __name__ == "__main__":
    import argparse
    workspace = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    parser = argparse.ArgumentParser(description="Swind Agent 01B: Đọc BOQ từ ảnh khách gửi")
    parser.add_argument("--image-dir", default=os.path.join(workspace, "input picture"), help="Thư mục chứa ảnh BOQ")
    parser.add_argument("--output-xlsx", default=None, help="Đường dẫn file Excel BOQ đầu ra (nếu gộp)")
    parser.add_argument("--merge-images", action="store_true", help="Gộp tất cả ảnh vào 1 file BOQ duy nhất (Mặc định: Tách mỗi ảnh 1 khách riêng)")
    parser.add_argument("--output-json", default=os.path.join(workspace, "data", "master_data", "last_image_boq_ocr.json"), help="File JSON log đối soát OCR")
    args = parser.parse_args()

    reader = Agent01B_ImageBOQReader(workspace)
    res = reader.process_image_folder(
        args.image_dir,
        args.output_xlsx,
        args.output_json,
        split_per_image=not args.merge_images,
    )
    if res["status"] == "SUCCESS":
        print("\n--- BẢNG ĐỐI SOÁT DỮ LIỆU ĐỌC TỪ ẢNH ---")
        for i, r in enumerate(res["rows"], 1):
            note_disp = f" | Ghi chú: [{r['note']}]" if r.get("note") else ""
            print(f"  {i:02d}. [{r['source_image'][:16]}...] {r['description']} | ĐVT: {r['unit']} | SL: {r['qty']}{note_disp}")
