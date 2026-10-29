import openpyxl
import re
import json
import os
import unicodedata
from swind_utils import detect_boq_header, get_subtype_rules, extract_duct_fitting_geometry

FILTER_PREFIXES = ("lọc bụi", "lọc thô", "lọc sơ cấp", "lọc g4", "phin lọc", "lưới lọc", "tấm lọc", "bộ lọc", "khung lọc", "filter g4", "g4 filter", "pre filter", "pre-filter")
FINE_FILTER_RX = re.compile(r'\b(?:f[5-9]|m[5-6]|e1[0-2]|h1[0-4]|hepa|ulpa)\b|than hoạt tính|tĩnh điện|túi|carbon|bag')


def is_standalone_coarse_filter(name_lower):
    """Dòng BOQ là lọc bụi thô / G4 đứng riêng (không phải lọc tinh F5-F9/HEPA, than hoạt tính, lọc tĩnh điện, lọc túi)."""
    n = re.sub(r'^\s*[-•*+]+\s*', '', name_lower)
    n = re.sub(r'^(?:cung cấp|gia công|sản xuất|lắp đặt)(?:\s*(?:,|và|&)\s*(?:cung cấp|gia công|sản xuất|lắp đặt))*\s*:?\s*', '', n)
    return n.startswith(FILTER_PREFIXES) and not FINE_FILTER_RX.search(n)


def _safe_qty(v):
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


class SwindBOQParser:
    """
    Trình phân tích BOQ tự động, nhận diện cấu trúc động:
    - Tìm dòng tiêu đề (header) động.
    - Phân biệt dòng phân nhóm (Category Header) và dòng sản phẩm thực tế.
    - Trích xuất thông số kích thước (W, H, L) và chủng loại sản phẩm.
    """
    def __init__(self, boq_filepath, wb=None, include_flex_duct=False):
        self.filepath = boq_filepath
        self.wb = wb if wb is not None else openpyxl.load_workbook(boq_filepath, data_only=True)
        self.include_flex_duct = bool(include_flex_duct)
        self.commercial_catalog = []
        self.manufacturable_catalog = []
        self.skipped_rows = {}
        self.missing_dim_rows = {}
        self.skip_reasons = {}
        self.parse_errors = {}
        self.row_overrides = {}
        self._load_catalogs()
        
        self.learned_keywords = {}
        learning_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data/master_data/boq_keyword_learning.json")
        if os.path.exists(learning_path):
            with open(learning_path, "r", encoding="utf-8") as f:
                self.learned_keywords = json.load(f)

    def _load_catalogs(self):
        try:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            candidates = [
                os.path.join(base_dir, "../data/master_data"),
                os.path.join(base_dir, "data/master_data"),
                os.path.abspath("data/master_data")
            ]
            comm_path, mfg_path = None, None
            for cand in candidates:
                cp = os.path.join(cand, "commercial_catalog.json")
                mp = os.path.join(cand, "manufacturable_catalog.json")
                if os.path.exists(cp) and os.path.exists(mp):
                    comm_path, mfg_path = cp, mp
                    break
            if comm_path and mfg_path:
                with open(comm_path, "r", encoding="utf-8") as f:
                    self.commercial_catalog = json.load(f)
                with open(mfg_path, "r", encoding="utf-8") as f:
                    self.manufacturable_catalog = json.load(f)
                print(f"[SwindBOQParser] Đã nạp thành công {len(self.commercial_catalog)} mục thương mại & {len(self.manufacturable_catalog)} mục sản xuất.")
            else:
                print(f"[SwindBOQParser] CẢNH BÁO: Không tìm thấy file catalog tại: {candidates}")
        except Exception as e:
            print(f"[SwindBOQParser] Lỗi khi nạp catalogs: {e}")

    def parse_sheet(self, sheet_name, max_items=None, _override_col_map=None, _override_header_row=None, _initial_section=None, _is_side_table=False):
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Sheet {sheet_name} không tồn tại trong {self.filepath}")

        if not _is_side_table and _override_col_map is None and _override_header_row is None and max_items is None and not self.row_overrides:
            cache_dict = getattr(self.wb, "_swind_parsed_cache", None)
            cache_key = (sheet_name, self.include_flex_duct)
            if cache_dict and cache_key in cache_dict:
                c_items, c_skip, c_miss, c_reasons = cache_dict[cache_key]
                self.skipped_rows[sheet_name] = set(c_skip)
                self.missing_dim_rows[sheet_name] = list(c_miss)
                self.skip_reasons[sheet_name] = dict(c_reasons)
                return [dict(it) for it in c_items]
            
        ws = self.wb[sheet_name]
        
        # 1. Nhận diện Header BOQ (dùng hàm chung từ swind_utils)
        header_row, col_map, hidden_rows, hidden_cols, max_scan_col = detect_boq_header(ws)
        if _override_col_map is not None:
            col_map = dict(_override_col_map)
        if _override_header_row is not None:
            header_row = _override_header_row
        combo_cfg = get_subtype_rules().get("COMBO_AND_DETECTION_RULES", {})
        cv_kws = tuple(combo_cfg.get("missing_qty_cv_keywords", ("louver", "luover", "louvre", "miệng gió", "cửa gió", "cửa nan", "van ", "vcd", "mfd", "mfsd", "nrd", "prd"))) + ("cửa chớp", "chớp nhôm")
        cv_rx = combo_cfg.get("missing_qty_cv_regex", r'\b(fd|md|sag|rag|eag|oag|pag|seg|eal|oal|fal)\b')

        image_rows = set()
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            if self.filepath and os.path.exists(self.filepath) and zipfile.is_zipfile(self.filepath):
                ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
                with zipfile.ZipFile(self.filepath, "r") as zf:
                    for zname in zf.namelist():
                        if "drawings/drawing" in zname and zname.endswith(".xml"):
                            root = ET.fromstring(zf.read(zname))
                            for anc in list(root.findall("xdr:oneCellAnchor", ns)) + list(root.findall("xdr:twoCellAnchor", ns)):
                                if anc.find("xdr:pic", ns) is not None:
                                    fr = anc.find("xdr:from", ns)
                                    if fr is not None and fr.find("xdr:row", ns) is not None:
                                        r_img = int(fr.find("xdr:row", ns).text) + 1
                                        if r_img > header_row:
                                            image_rows.add(r_img)
        except Exception:
            pass
            
        items = []
        if not _is_side_table:
            self.skipped_rows[sheet_name] = set()
            self.missing_dim_rows[sheet_name] = []
            self.skip_reasons[sheet_name] = {}
        current_section = _initial_section if _initial_section else "SẢN PHẨM"
        current_section_aux = ""
        empty_streak = 0
        sheet_default_flange = ""
        for ws_chk in [ws] + [self.wb[sn] for sn in self.wb.sheetnames if sn != sheet_name]:
            for hr_scan in range(1, min(20, ws_chk.max_row + 1)):
                for hc_scan in range(1, min(25, ws_chk.max_column + 1)):
                    hv = unicodedata.normalize('NFC', str(ws_chk.cell(hr_scan, hc_scan).value or "")).lower()
                    if "bích c" in hv or "nẹp c" in hv:
                        sheet_default_flange = "NEP_C"
                        break
                if sheet_default_flange:
                    break
            if sheet_default_flange:
                break
        
        for r in range(header_row + 1, ws.max_row + 1):
            if r in hidden_rows:
                continue
                
            name_val = ws.cell(r, col_map.get("name", 2)).value
            qty_val = ws.cell(r, col_map["qty"]).value if "qty" in col_map else None
            unit_val = ws.cell(r, col_map["unit"]).value if "unit" in col_map else None
            col1_val = ws.cell(r, 1).value if (col_map.get("name", 2) > 1 and "stt" not in col_map) else None
            col1_str = unicodedata.normalize('NFC', str(col1_val).strip()) if isinstance(col1_val, str) and len(str(col1_val).strip()) > 2 and not re.match(r'^(?:[\dIVXLC]+(?:[\.\-]\d+)*\.?|CT\d+)$', str(col1_val).strip(), re.IGNORECASE) else ""
            if not name_val and col1_str and qty_val is None and unit_val is None:
                name_val = col1_str
            
            if not name_val:
                if qty_val is None and unit_val is None:
                    empty_streak += 1
                    if empty_streak >= 40:
                        break
                continue
            empty_streak = 0
                
            name_str = unicodedata.normalize('NFC', str(name_val).strip())
            if (name_str.lower() == "nội dung công việc" or any(k in name_str.lower() for k in ("đại diện ban", "đại diện chủ đầu tư", "đại diện nhà thầu"))) and len(items) > 0:
                break
            raw_name_str = name_str
            is_missing_qty = False
            is_pure_dim = bool(re.match(r'^(?:[dD]\d+|\d+\s*[xX*]\s*\d+)$', name_str))
            if name_str.lower().endswith(("kích thước:", "kích thước")) and not re.search(r'\d', name_str):
                current_section = name_str
                continue
            if is_pure_dim and current_section and current_section.upper() != "SẢN PHẨM":
                name_str = f"{current_section.strip()} {name_str}"
            
            # Kiểm tra xem có phải dòng tiêu đề phân nhóm không (giữ lại dòng Cửa/Van có quy cách kích thước dù chưa điền khối lượng)
            if qty_val is None or not isinstance(qty_val, (int, float)) or qty_val <= 0:
                spec_chk = str(ws.cell(r, col_map["spec"]).value or "").strip() if "spec" in col_map else ""
                model_chk = str(ws.cell(r, col_map["model"]).value or "").strip() if "model" in col_map else ""
                nl_chk = f"{name_str} {spec_chk} {model_chk}".lower()
                is_blank_qty = (qty_val is None or str(qty_val).strip() == "")
                is_zero_qty_with_unit = (not col_map.get("is_order_sheet") and qty_val in (0, 0.0, "0") and unit_val not in (None, "", 0, "0"))
                has_cv_kw = any(k in nl_chk for k in cv_kws) or bool(re.search(cv_rx, nl_chk))
                has_cv_dim = bool(re.search(r'([wWaAbB]?\s*\d+\s*[xX*]\s*[hHbBwW]?\s*\d+|\b[dD][nN]?\s*\d{2,4}\b|\b(?:vcd|fd|nrd|prd|md)\s+\d{2,4}\b)', nl_chk))
                if (is_zero_qty_with_unit or (is_blank_qty and not col_map.get("is_order_sheet") and has_cv_kw and has_cv_dim)) and "tổng cộng" not in nl_chk:
                    if "qty" in col_map:
                        is_missing_qty = is_blank_qty
                        qty_val = 0.0
                    else:
                        qty_val = 1.0
                else:
                    is_valid_unit_str = isinstance(unit_val, str) and unit_val.strip() != "" and not unit_val.strip().replace('.', '').isdigit()
                    is_parent_title_kw = (not has_cv_dim) and any(k in nl_chk for k in ("cửa", "miệng", "van", "ống", "louver", "luover", "louvre", "grille", "diffuser", "damper", "khớp nối", "tiêu âm"))
                    if (not is_valid_unit_str or is_parent_title_kw) and not is_pure_dim and "tổng cộng" not in nl_chk:
                        if len(name_str) > 2:
                            current_section = name_str
                            current_section_aux = col1_str if col1_str.lower() != name_str.lower() else ""
                    continue
                
            if not unit_val:
                if "stt" in col_map and col_map["stt"] > 0 and ws.cell(r, col_map["stt"]).value in (0, "0") and not re.search(r'\d', name_str):
                    current_section = name_str
                    continue
                unit_val = "m" if any(k in name_str.lower() for k in ("ống đồng", "ông đồng", "cách nhiệt", "upvc")) else "Cái"
                
            # Đọc kích thước (Chỉ đọc từ cột tương ứng nếu header thực sự có cột đó)
            w_val = ws.cell(r, col_map["w"]).value if "w" in col_map else None
            h_val = ws.cell(r, col_map["h"]).value if "h" in col_map else None
            w2_val = ws.cell(r, col_map["w2"]).value if "w2" in col_map else None
            h2_val = ws.cell(r, col_map["h2"]).value if "h2" in col_map else None
            l_val = ws.cell(r, col_map["l"]).value if "l" in col_map else None
            r_or_d_val = ws.cell(r, col_map["r_or_d"]).value if "r_or_d" in col_map else None
            offset_col_val = ws.cell(r, col_map["offset"]).value if "offset" in col_map else None
            thk_val = ws.cell(r, col_map["thickness"]).value if "thickness" in col_map else None
            flange_col_val = ws.cell(r, col_map["flange"]).value if "flange" in col_map else None
            flange_col_str = unicodedata.normalize('NFC', str(flange_col_val).strip()).lower() if flange_col_val not in (None, 0, "0", "#REF!", "#VALUE!", "#N/A") else ""
            spec_val = ws.cell(r, col_map["spec"]).value if "spec" in col_map else None
            spec_str = unicodedata.normalize('NFC', str(spec_val).strip()) if spec_val not in (None, 0, "0", "#REF!", "#VALUE!", "#N/A") else ""
            # Fallback: đọc cột kề phải cột name nếu col_map thiếu cột kích thước (BOQ không có header row)
            if not spec_str and "spec" not in col_map and "model" not in col_map and "w" not in col_map:
                fallback_col = col_map.get("name", 2) + 1
                mapped_cols = {v for k, v in col_map.items() if isinstance(v, int)}
                if fallback_col not in mapped_cols:
                    fb_val = ws.cell(r, fallback_col).value
                    if fb_val is not None:
                        fb_str = unicodedata.normalize('NFC', str(fb_val).strip())
                        if re.search(r'\d{2,4}\s*[xX×*]\s*\d{2,4}', fb_str) or (isinstance(fb_val, str) and len(fb_str) > 2 and fb_str.lower() != name_str.lower()):
                            spec_str = fb_str
            brand_col_val = ws.cell(r, col_map["brand"]).value if "brand" in col_map else None
            brand_col_str = unicodedata.normalize('NFC', str(brand_col_val).strip()) if brand_col_val not in (None, 0, "0", "#REF!", "#VALUE!", "#N/A") else ""
            model_val = ws.cell(r, col_map["model"]).value if "model" in col_map else None
            if model_val in (0, "0", "#REF!", "#VALUE!", "#N/A"):
                model_val = None
            remarks_val = ws.cell(r, col_map["remarks"]).value if "remarks" in col_map else None
            note_boq_raw = unicodedata.normalize('NFC', str(remarks_val).strip()) if remarks_val not in (None, 0, "0") else ""
            extra_spec = brand_col_str if any(k in brand_col_str.lower() for k in ("tôn", "tole", "dày", "day", "0.", "1.", "mm", "hoa sen", "phương nam", "inox")) else ""
            flange_note = "Nẹp C" if flange_col_str in ("c", "nẹp c", "nep c", "bích c", "bich c") else (flange_col_str if flange_col_str else ("Nẹp C" if sheet_default_flange == "NEP_C" else ""))
            note_boq = " ".join(x for x in (spec_str, extra_spec, flange_note, note_boq_raw) if x)
            
            w = float(w_val) if isinstance(w_val, (int, float)) else 0.0
            h = float(h_val) if isinstance(h_val, (int, float)) else 0.0
            w2 = float(w2_val) if isinstance(w2_val, (int, float)) else 0.0
            h2 = float(h2_val) if isinstance(h2_val, (int, float)) else 0.0
            l = float(l_val) if isinstance(l_val, (int, float)) else 0.0
            r_or_d_num = float(r_or_d_val) if isinstance(r_or_d_val, (int, float)) else 0.0
            offset_num = float(offset_col_val) if isinstance(offset_col_val, (int, float)) else 0.0
            if isinstance(thk_val, (int, float)):
                thickness = float(thk_val)
            elif isinstance(thk_val, str):
                m_thk_col = re.search(r'(\d+[.,]\d+)', thk_val)
                thickness = float(m_thk_col.group(1).replace(',', '.')) if m_thk_col else 0.0
            else:
                thickness = 0.0
            
            # Chuẩn bị tên đầy đủ kèm tiền tố danh mục cha (section)
            clean_section = current_section.strip()
            full_name_str = name_str
            if clean_section and clean_section.upper() != "SẢN PHẨM" and clean_section.lower() not in name_str.lower():
                clean_section = re.sub(r'^[\dIVXLC]+\.\s*', '', clean_section)
                full_name_str = f"{clean_section} {name_str}"
                
            # Phân loại sản phẩm (Category detection) động từ Master Catalog
            aux_context = " ".join(x for x in (current_section_aux, col1_str if col1_str.lower() != raw_name_str.lower() else "") if x).strip()
            name_lower = name_str.lower().replace('ồ', 'ố')
            full_name_lower = f"{full_name_str} {aux_context}".lower().strip()
            sec_lower = f"{current_section} {current_section_aux}".lower().replace('ồ', 'ố').strip()
            model_str = str(model_val).strip() if model_val is not None else ""
            category = "KHAC"
            
            # Hàm kiểm tra từ khóa an toàn (chống va chạm từ con, ví dụ: 'chống cháy' chứa 'ống')
            def match_keyword(kw, text):
                if kw not in text:
                    return False
                if " " in kw:
                    return True
                return bool(re.search(r'\b' + re.escape(kw) + r'\b', text))

            # Kiểm tra loại trừ hệ ống cấp thoát nước / ống chữa cháy / ống đồng / uPVC / PPR / thép đen / thép tráng kẽm
            plumbing_kws = (
                "thép hàn", "thép ren", "thép đen", "thép tráng kẽm", "thép mạ kẽm", "ống thép",
                "u.pvc", "upvc", "ppr", "hdpe", "ống nhựa", "ống đồng", "ông đồng",
                "ống bảo ôn", "bảo ôn ống đồng", "gen cách nhiệt", "superlon",
                "van cổng", "van bướm", "van bi", "van giảm áp", "van y lọc", "van an toàn", "van xả khí",
                "van một chiều lá lật", "van một chiều lò xo", "van báo động", "van tràn ngập",
                "đầu phun", "sprinkler", "bình chữa cháy", "cuộn vòi", "tủ chữa cháy", "trụ cứu hỏa",
                "họng tiếp nước", "khớp nối mềm cao su", "khớp nối mềm inox", "rọ bơm", "y lọc"
            )
            is_hvac_prd = ("van giảm áp" in full_name_lower or "van xả áp" in full_name_lower) and (
                bool(re.search(r'\b\d{2,4}\s*[xX*]\s*\d{2,4}\b', f"{full_name_str} {note_boq}"))
                or any(k in f"{full_name_lower} {note_boq.lower()}" for k in ("tôn", "tole", "ei", "chịu lửa", "chống cháy", "mfd", "động cơ", "on/off", "on-off"))
            )
            is_plumbing_item = (not is_hvac_prd) and any(pk in name_lower or pk in full_name_lower for pk in plumbing_kws)
            if not is_plumbing_item and any(sk in sec_lower for sk in ("chữa cháy", "cấp nước", "thoát nước", "ống đồng", "bảo ôn ống")):
                if not any(dk in name_lower for dk in ("ống gió", "tôn", "tole", "cửa gió", "miệng gió", "van gió", "vcd", "mfd", "fd", "nrd", "prd", "obd", "louver")):
                    is_plumbing_item = True
            if not is_plumbing_item and name_lower.startswith(("cút ", "tê ", "y ", "chếch ", "côn ", "măng sông", "rắc co", "bích thép", "nút bịt")):
                if re.search(r'\b(?:[dD][nN]?\s*(?:15|20|21|25|27|32|34|40|42|48|50|60|65|75|80|90|110|114|140|160|168|219)|PN\s*\d+|SCH\s*\d+)\b', name_str, re.IGNORECASE):
                    if not any(dk in name_lower for dk in ("ống gió", "tôn", "tole")):
                        is_plumbing_item = True

            # 1. Kiểm tra Học máy (Learned keywords)
            if name_lower in self.learned_keywords:
                category = self.learned_keywords[name_lower]
            elif full_name_lower in self.learned_keywords:
                category = self.learned_keywords[full_name_lower]
            elif is_plumbing_item:
                category = "THUONG_MAI"
            else:
                # 2. Kiểm tra các trường hợp đặc thù ưu tiên cao (tránh bị từ khóa ngắn như 'van', 'cửa', 'miệng', 'fd' bắt nhầm)
                # 2a. Van nước / Cửa sổ / Thiết bị bắt đầu bằng Quạt, Máy lạnh, Dàn lạnh, Dàn nóng -> THUONG_MAI
                if name_lower.startswith("cửa sổ") and ("cấp bù gió" in name_lower or "khi có cháy" in name_lower):
                    category = "VAN_GIO"
                elif is_standalone_coarse_filter(name_lower):
                    # Lọc thô G4 đứng riêng (có kích thước) -> Swind sản xuất: báo giá theo công thức FILTER_G4 (Agent 06)
                    category = "CUA_GIO"
                elif name_lower.startswith(("quạt ", "máy lạnh", "dàn lạnh", "dàn nóng", "hướng dòng", "thiết bị xử lý", "cửa sổ", "tủ điện", "tủ điều khiển", "đầu cảm biến", "đầu báo", "cảm biến", "dây tín hiệu", "cáp điều khiển", "dây điện", "khối lượng sắt thép", "giá đỡ ống", "quang treo", "hồ sơ kiểm định", "chi phí ", "chi phi ", "vật tư phụ", "lượng tác nhân lạnh", "phụ kiện ống nước ngưng", "lò xo giảm chấn", "cao su giảm chấn", "chụp hút bếp", "điều hòa cục bộ", "+ điều hòa cục bộ")):
                    category = "THUONG_MAI"
                else:
                    for comm in self.commercial_catalog:
                        if comm.get("product_code") in ("VALVE_WATER", "PIPE_COPPER"):
                            if comm.get("product_code") == "VALVE_WATER" and (is_hvac_prd or any(k in name_lower for k in ("lửa", "khói", "cháy", "mfd", "mfsd", "fd", "vcd", "nrd", "prd", "md", "xả áp"))):
                                continue
                            if comm.get("product_code") == "PIPE_COPPER" and any(k in name_lower for k in ("louver", "louvre", "cửa gió", "miệng gió")):
                                continue
                            if any(match_keyword(kw, name_lower) for kw in comm.get("keywords", [])):
                                category = "THUONG_MAI"
                                break

                # 2b. Box / Hộp tĩnh áp tôn mạ kẽm hoặc Phụ kiện ống gió đầu dòng (Cút, Côn, Tê, Bịt ống, Nối ống...) -> ONG_GIO
                if category == "KHAC" and name_lower.startswith(("box ", "hộp ", "hôp ")) and not name_lower.startswith(("hộp và cửa", "hộp + cửa", "hộp kèm cửa", "box và cửa", "box + cửa", "hộp quạt", "hộp điện", "hộp nối", "hộp cứu hỏa", "hộp chữa cháy")) and "tiêu âm" not in name_lower:
                    category = "ONG_GIO"
                if category == "KHAC" and not is_plumbing_item and "tiêu âm" not in name_lower and name_lower.startswith(("cút ", "co ", "chếch ", "côn ", "tê ", "chạc ", "chân rẽ", "gót giày", "cổ trích", "bịt đầu", "bịt ống", "nắp bịt", "măng xông", "măng sông", "nối ống", "zét ", "ống lượn")):
                    if re.search(r'\b(?:[dD][nN]?|Ø|phi\s*)\d{2,4}\b|\b\d{2,4}\s*[xX*]\s*\d{2,4}\b', name_str, re.IGNORECASE):
                        category = "ONG_GIO"

                if category == "KHAC" and any(k in name_lower or k in note_boq.lower() for k in ("access door", "cửa công tác", "cửa thăm", "cửa kiểm tra", "cửa chớp", "chớp nhôm", "cửa bầu dục", "cửa nan z", "cửa sọt trứng", "cửa khuếch tán")):
                    category = "CUA_GIO"
                if category == "KHAC" and name_lower.startswith(("cửa gió", "miệng gió", "cửa cấp", "cửa hồi", "cửa hút", "cửa thải", "cửa lấy", "cửa xả", "cửa tươi", "cửa sọt", "cửa nan", "cửa khuếch tán", "cửa bầu dục", "cửa louver", "cửa chớp", "louver ")):
                    category = "CUA_GIO"
                if category == "KHAC" and any(k in sec_lower for k in ("cửa gió", "miệng gió", "cửa chớp", "chớp nhôm", "lưới chắn chim", "louver", "luover", "louvre", "grille", "griller", "diffuser", "insect screen", "i.s", "lcct", "lưới chắn côn trùng", "fag", "sag", "rag", "eag", "seag", "kag", "mag", "fal", "eal", "oal")):
                    if not name_lower.startswith(("van ", "damper", "volume control", "fire damper", "motorized damper", "non-return", "ống ", "cút ", "côn ", "tê ", "chân rẽ", "hộp ", "chụp hút", "chụp thông")) and not any(k in sec_lower for k in ("chụp hút", "chụp thông")) and not any(k in name_lower for k in ("volume control damper", "fire damper", "motorized damper", "non-return damper")):
                        category = "CUA_GIO"

                # 2c. Kiểm tra Hàng Sản Xuất có độ đặc thù cao (Canvas, Chụp thông hơi, Tiêu âm, Cửa, Van, Ống mềm)
                if category == "KHAC":
                    for mfg in self.manufacturable_catalog:
                        if mfg.get("mapping_category") not in ("ONG_GIO", "PHU_KIEN_ONG", "KHAC"):
                            if any(match_keyword(kw, name_lower) for kw in mfg.get("keywords", [])):
                                category = mfg.get("mapping_category")
                                break
                    if category == "KHAC" and (note_boq or model_str):
                        note_lower_cat = f"{note_boq} {model_str}".lower()
                        for mfg in self.manufacturable_catalog:
                            if mfg.get("mapping_category") not in ("ONG_GIO", "PHU_KIEN_ONG", "KHAC"):
                                if any(match_keyword(kw, note_lower_cat) for kw in mfg.get("keywords", [])):
                                    category = mfg.get("mapping_category")
                                    break
                    if category == "KHAC" and full_name_lower != name_lower:
                        for mfg in self.manufacturable_catalog:
                            if mfg.get("mapping_category") not in ("ONG_GIO", "PHU_KIEN_ONG", "KHAC"):
                                if any(match_keyword(kw, full_name_lower) for kw in mfg.get("keywords", [])):
                                    category = mfg.get("mapping_category")
                                    break
    
                # 3. Nếu chưa khớp, kiểm tra toàn bộ Hàng Thương Mại còn lại
                if category == "KHAC":
                    for comm in self.commercial_catalog:
                        if any(match_keyword(kw, name_lower) for kw in comm.get("keywords", [])):
                            category = "THUONG_MAI"
                            break
                    if category == "KHAC" and full_name_lower != name_lower:
                        for comm in self.commercial_catalog:
                            if any(match_keyword(kw, full_name_lower) for kw in comm.get("keywords", [])):
                                category = "THUONG_MAI"
                                break
                        
                # 4. Nếu vẫn chưa khớp, kiểm tra Ống Gió & Phụ kiện ống gió
                if category == "KHAC":
                    for mfg in self.manufacturable_catalog:
                        if mfg.get("mapping_category") in ("ONG_GIO", "PHU_KIEN_ONG"):
                            if any(match_keyword(kw, name_lower) for kw in mfg.get("keywords", [])):
                                category = "ONG_GIO"
                                break
                    if category == "KHAC" and full_name_lower != name_lower:
                        for mfg in self.manufacturable_catalog:
                            if mfg.get("mapping_category") in ("ONG_GIO", "PHU_KIEN_ONG"):
                                if any(match_keyword(kw, full_name_lower) for kw in mfg.get("keywords", [])):
                                    category = "ONG_GIO"
                                    break

            if category == "PHU_KIEN_ONG":
                category = "ONG_GIO"
            elif category in ("KHOP_NOI_MEM", "CHUP_THONG_HOI"):
                category = "CANVAS"

            # 5. Fallback heuristics có kiểm soát chặt chẽ
            if category == "KHAC" and any(k in name_lower for k in ("kẹp bích", "kep bich", "ty treo", "bulong", "bu lông")):
                category = "THUONG_MAI"
            elif category == "KHAC" and not is_plumbing_item:
                has_duct = bool(re.search(r'\bống\s+gió\b|^ống\s+(?:kt\b|kích\s*thước\b|\d{2,4}\s*[xX*]\s*\d{2,4})|\btole\b|\btôn\b|(?:^|\s)z\s*\d{2,4}\s*[xX*]\s*\d{2,4}|(?:^|\s)z\s*chuyển|\b\d{2,4}\s*[xX*]\s*\d{2,4}\s*(?:[/\-xX*\s]+)?(?:[dD][nN]?|Ø|phi)\s*\d{2,4}\b', name_lower))
                if has_duct or any(k in sec_lower for k in ("ống gió", "phụ kiện ống gió", "côn, cút", "hút mùi bếp", "hút mùi wc", "thông gió")):
                    category = "ONG_GIO"
                
            # Trích xuất kích thước: Ưu tiên tìm trên tên sản phẩm + model trước khi mở rộng sang cột Ghi chú BOQ
            primary_search_text = f"{raw_name_str} {model_str}".strip().replace('×', 'x').replace('，', ',').replace('：', ':')
            primary_clean_text = re.sub(r'\b\d+(?:[.,]\d+)?\s*(?:kW|W|HP|l/s|m3/h|Pa|V|Hz|dB|A)\b', ' ', primary_search_text, flags=re.IGNORECASE)
            dim_search_text = f"{name_str} {model_str} {note_boq}".strip().replace('×', 'x').replace('，', ',').replace('：', ':')
            # Loại bỏ các thông số công suất/lưu lượng/áp suất để tránh bắt nhầm
            dim_clean_text = re.sub(r'\b\d+(?:[.,]\d+)?\s*(?:kW|W|HP|l/s|m3/h|Pa|V|Hz|dB|A)\b', ' ', dim_search_text, flags=re.IGNORECASE)
            d_val = 0.0
            fitting_geom = None

            if category in ("ONG_GIO", "CANVAS", "TIEU_AM"):
                fitting_geom = extract_duct_fitting_geometry(dim_search_text, default_w=w, default_h=h, default_d=d_val, default_w2=w2, default_h2=h2)
                if (w == 0 or h == 0) and fitting_geom["w1"] > 0 and fitting_geom["h1"] > 0:
                    w = fitting_geom["w1"]
                    h = fitting_geom["h1"]
                if (w2 == 0 or h2 == 0) and fitting_geom["w2"] > 0 and fitting_geom["h2"] > 0:
                    w2 = fitting_geom["w2"]
                    h2 = fitting_geom["h2"]
                if d_val == 0 and fitting_geom["d1"] > 0:
                    d_val = fitting_geom["d1"]
                    if w == 0 and h == 0:
                        w = d_val
                        h = d_val
                if r_or_d_num > 0:
                    if any(k in name_lower for k in ("vuông tròn", "vuong tron", "côn", "con ", "chân rẽ", "chan re", "cổ trích", "hộp gió", "hop gio")) and d_val == 0:
                        d_val = r_or_d_num
                        fitting_geom["d1"] = r_or_d_num
                    elif fitting_geom["radius"] == 0:
                        fitting_geom["radius"] = r_or_d_num
                if offset_num > 0 and fitting_geom["offset_h"] == 0:
                    fitting_geom["offset_h"] = offset_num
                if l == 0 and fitting_geom["length"] > 0:
                    l = fitting_geom["length"]
                if thickness == 0 and fitting_geom["thickness"] > 0:
                    thickness = fitting_geom["thickness"]
                if flange_note:
                    fl_up = flange_note.upper()
                    if "C" in fl_up and "TDC" not in fl_up:
                        fitting_geom["flange_type"] = "NEP_C"
                    elif "TDC" in fl_up or "BÍCH" in fl_up:
                        fitting_geom["flange_type"] = "TDC"

            l_explicit = re.search(r'(?:\bL\s*[:=]?\s*|\bdài\s+)(\d{2,4})\b', primary_clean_text, re.IGNORECASE) or re.search(r'(?:\bL\s*[:=]?\s*|\bdài\s+)(\d{2,4})\b', dim_clean_text, re.IGNORECASE)
            if l_explicit and l == 0:
                l = float(l_explicit.group(1))

            # Bóc tách đường kính tròn D / DN / Ø / Phi (kể cả viết dạng D100x100: "cứ kích thước D là tròn")
            d_match = re.search(r'(?:^|[\s_(+:-])(?:[dD][nN]?|Ø|phi\s*)(\d{2,4})(?:\s*[xX*]\s*\1)?(?:\s*mm)?(?:\b|(?=[^\w]))', primary_clean_text, re.IGNORECASE) or re.search(r'(?:^|[\s_(+:-])(?:[dD][nN]?|Ø|phi\s*)(\d{2,4})(?:\s*[xX*]\s*\1)?(?:\s*mm)?(?:\b|(?=[^\w]))', dim_clean_text, re.IGNORECASE)
            rect_primary_text = re.sub(r'(?<![a-zA-Z])(?:[dD][nN]?|Ø|phi\s*)(\d{2,4})\s*[xX*]\s*\1\b', ' ', primary_clean_text, flags=re.IGNORECASE)
            rect_primary_text = re.sub(r'\b(?:to|mắt\s*lưới|lưới)\s*\d{1,2}\s*[xX*]\s*\d{1,2}\b|\b[vV]\s*\d{2,3}\s*[xX*]\s*\d{2,3}(?:\s*[xX*]\s*\d{1,2})?\b', ' ', rect_primary_text, flags=re.IGNORECASE)
            rect_search_text = re.sub(r'(?<![a-zA-Z])(?:[dD][nN]?|Ø|phi\s*)(\d{2,4})\s*[xX*]\s*\1\b', ' ', dim_clean_text, flags=re.IGNORECASE)
            rect_search_text = re.sub(r'\b(?:to|mắt\s*lưới|lưới)\s*\d{1,2}\s*[xX*]\s*\d{1,2}\b|\b[vV]\s*\d{2,3}\s*[xX*]\s*\d{2,3}(?:\s*[xX*]\s*\d{1,2})?\b', ' ', rect_search_text, flags=re.IGNORECASE)

            w_neck, h_neck = 0.0, 0.0
            has_explicit_kt_bi = False
            kt_bi_rx = r'(?:kt|kích\s*thước)\s*(?:phủ\s*)?bì(?:(?!kt\s*c[ổôo]|kích\s*thước\s*c[ổôo]|\bktc\b).)*?(\d{2,4})\s*[xX*]\s*(\d{2,4})'
            kt_co_rx = r'(?:(?:kt|kích\s*thước)\s*c[ổôo]|\bktc\b)[^0-9]*?(\d{2,4})\s*[xX*]\s*(\d{2,4})'
            m_bi = re.search(kt_bi_rx, note_boq, re.IGNORECASE) or re.search(kt_bi_rx, name_str, re.IGNORECASE)
            m_co = re.search(kt_co_rx, note_boq, re.IGNORECASE) or re.search(kt_co_rx, name_str, re.IGNORECASE)
            if m_co:
                w_neck = float(m_co.group(1))
                h_neck = float(m_co.group(2))
            if m_bi and w == 0 and h == 0:
                w = float(m_bi.group(1))
                h = float(m_bi.group(2))
                has_explicit_kt_bi = True

            if w == 0 and h == 0:
                dim_3d = re.search(r'[wWaAbB]?\s*(\d+)\s*[xX*]\s*[hHbBwW]?\s*(\d+)\s*[xX*]\s*[lLdD]?\s*(\d+)', rect_primary_text)
                dim_2d = re.search(r'[wWaAbB]?\s*(\d+)\s*[xX*]\s*[hHbBwW]?\s*(\d+)', rect_primary_text)
                if not dim_3d and not dim_2d:
                    dim_3d = re.search(r'[wWaAbB]?\s*(\d+)\s*[xX*]\s*[hHbBwW]?\s*(\d+)\s*[xX*]\s*[lLdD]?\s*(\d+)', rect_search_text)
                    dim_2d = re.search(r'[wWaAbB]?\s*(\d+)\s*[xX*]\s*[hHbBwW]?\s*(\d+)', rect_search_text)
                # Hỗ trợ bắt lỗi gõ nhầm chữ 'x' thành số '6' dạng 10006300 (4 chữ số + 6 + 3 chữ số tròn trăm)
                dim_typo = re.search(r'\b(\d{3,4})6(\d00)\b', rect_search_text)
                if dim_3d and not (category == "CUA_GIO" and dim_2d and dim_2d.start() < dim_3d.start()):
                    w = float(dim_3d.group(1))
                    h = float(dim_3d.group(2))
                    l = float(dim_3d.group(3))
                elif dim_2d:
                    w = float(dim_2d.group(1))
                    h = float(dim_2d.group(2))
                    if dim_3d and l == 0:
                        l = float(dim_3d.group(3))
                    if category == "CUA_GIO" and w >= 100 and h == 0 and dim_2d.group(2) == "00":
                        h = w
                elif dim_typo:
                    w = float(dim_typo.group(1))
                    h = float(dim_typo.group(2))
                elif category == "CUA_GIO":
                    m_rect_d = re.search(r'\b(\d{3,4})\s*mm\s*[,;-]?\s*[dD](\d{2,3})\b', dim_clean_text, re.IGNORECASE)
                    m_lb = re.search(r'\bLB\s*(\d{2,3})\b', dim_clean_text, re.IGNORECASE)
                    m_sl = re.search(r'\b(?:SL|LS)\s*(\d)\b', dim_clean_text, re.IGNORECASE)
                    m_len = re.search(r'\b(\d{3,4})\s*mm\b', dim_clean_text, re.IGNORECASE)
                    if m_rect_d:
                        w = float(m_rect_d.group(1))
                        h = float(m_rect_d.group(2))
                        d_match = None
                    elif m_lb and m_len:
                        w = float(m_len.group(1))
                        h = float(m_lb.group(1))
                    elif m_sl and m_len:
                        n_slots = int(m_sl.group(1))
                        if 1 <= n_slots <= 6:
                            w = float(m_len.group(1))
                            h = float(50 * n_slots + 50)

            if d_match and w == 0 and h == 0:
                parsed_d = float(d_match.group(1))
                if 50 <= parsed_d <= 3000:
                    d_val = parsed_d
                    if category in ("VAN_GIO", "CUA_GIO", "ONG_GIO", "TIEU_AM"):
                        w = d_val
                        h = d_val

            # Bóc tách diện tích m2 (ví dụ: Cửa sổ cấp bù gió khi có cháy 1.5m2 / 0.8m2 / 0.75m2) hoặc 1 kích thước đơn (mặc định tròn D, ví dụ VCD 150)
            if w == 0 and h == 0 and category in ("VAN_GIO", "CUA_GIO"):
                area_match = re.search(r'(\d+(?:[.,]\d+)?)\s*m2\b', dim_search_text, re.IGNORECASE)
                if area_match:
                    try:
                        area_m2 = float(area_match.group(1).replace(',', '.'))
                        if area_m2 >= 20.0:
                            w = 10000.0
                            h = round((area_m2 / 10.0) * 1000.0)
                        elif area_m2 >= 10.0:
                            w = round((area_m2 / 2.0) * 1000.0)
                            h = 2000.0
                        elif area_m2 >= 1.0:
                            w = round(area_m2 * 1000.0)
                            h = 1000.0
                        elif area_m2 > 0:
                            w = 1000.0
                            h = round(area_m2 * 1000.0)
                    except:
                        pass
                elif d_val == 0:
                    no_ei_text = re.sub(r'(?<![a-zA-Z])EI\s*[-:]?\s*\d{2,3}\b', ' ', dim_clean_text, flags=re.IGNORECASE)
                    single_d_rx = combo_cfg.get("single_dim_damper_regex", r'\b(?:vcd|fd|nrd|prd|md|obd|mfd|mfsd)\s+(\d{2,4})\b')
                    single_d_match = re.search(single_d_rx, no_ei_text, re.IGNORECASE)
                    if single_d_match:
                        parsed_sd = float(single_d_match.group(1))
                        if 50 <= parsed_sd <= 2000:
                            d_val = parsed_sd
                            w = d_val
                            h = d_val

            # Ghi đè kích thước theo yêu cầu User (--override "Sheet!Dòng:w=..,h=..") trước khi xét thiếu kích thước
            row_ov = self.row_overrides.get((sheet_name, r)) if self.row_overrides else None
            if row_ov:
                for k_ov in ("w", "h", "d"):
                    if k_ov in row_ov:
                        try:
                            v_ov = float(row_ov[k_ov])
                        except (TypeError, ValueError):
                            continue
                        if k_ov == "w":
                            w = v_ov
                        elif k_ov == "h":
                            h = v_ov
                        else:
                            d_val = v_ov
            is_missing_dims = False
            orig_category = category
            if category in ("VAN_GIO", "CUA_GIO", "CANVAS") and w == 0 and h == 0 and d_val == 0:
                is_missing_dims = True
                category = "THUONG_MAI"
            elif category == "CUA_GIO" and d_val > 0 and any(k in full_name_lower for k in ("ead",)):
                category = "THUONG_MAI"
            if category == "ONG_GIO" and w == 0 and h == 0 and d_val == 0 and ("kẹp bích" in name_lower or (name_lower.startswith(("box ", "hộp ", "hôp ")) and str(unit_val).strip().lower() not in ("m2", "m²"))):
                category = "THUONG_MAI"
            if any(k in name_lower for k in ("măng xông", "mang xong", "măng sông", "mang song")):
                unit_val = "cái"

            # Trích xuất chỉ số chống cháy EI (EI30, EI60, EI90, EI120) — quét cả dòng con và tiêu đề cha
            ei_primary_text = f"{full_name_str} {model_str} {note_boq_raw}"
            ei_spec_text = "" if any(k in str(spec_str or "").lower() for k in ("smd", "không kiểm định")) else str(spec_str or "")
            full_dim_search_text = f"{full_name_str} {model_str} {note_boq}"
            ei_rating = 0
            ei_match = re.search(r'(?<![a-zA-Z])EI\s*[-:]?\s*(\d{2,3})\b', ei_primary_text, re.IGNORECASE) or re.search(r'(?<![a-zA-Z])EI\s*[-:]?\s*(\d{2,3})\b', ei_spec_text, re.IGNORECASE)
            if ei_match:
                ei_rating = int(ei_match.group(1))
            elif "cấp bù gió khi có cháy" in full_name_lower:
                ei_rating = 120

            # Trích xuất độ dày từ tên sản phẩm, tiêu đề cha hoặc model nếu có chỉ định rõ
            thk_match = re.search(r'(?:tole|tôn|zam|dày|day|thk|độ dày)\s*(\d+[.,]\d+)\s*mm', dim_search_text, re.IGNORECASE) or re.search(r'(?:tole|tôn|zam|dày|day|thk|độ dày)\s*(\d+[.,]\d+)\s*mm', full_dim_search_text, re.IGNORECASE)
            if not thk_match:
                thk_match = re.search(r'(\d+[.,]\d+)\s*mm', dim_search_text, re.IGNORECASE) or re.search(r'(\d+[.,]\d+)\s*mm', full_dim_search_text, re.IGNORECASE)
            if thk_match:
                try:
                    parsed_thk = float(thk_match.group(1).replace(',', '.'))
                    if 0.3 <= parsed_thk <= 5.0: # Giới hạn dải độ dày hợp lý
                        thickness = parsed_thk
                except:
                    pass
            elif category == "ONG_GIO" and "fcu" in name_lower and name_lower.startswith(("hộp ", "box ")):
                thickness = 0.75
                    
            name_str = full_name_str

            # Rule 1: Kích thước cửa gió ghi KTC hoặc cửa bù khí (mặc định nan Z) thì cộng thêm mỗi cạnh 50mm
            ktc_rx = combo_cfg.get("ktc_regex", r'\bktc\b|kích\s*thước\s*c[ổôo]|\bkt\s*c[ổôo]\b')
            raw_makeup_chk = f"{raw_name_str} {model_str} {note_boq}".lower()
            is_makeup_air_z = (
                any(k in raw_makeup_chk for k in ("cấp khí bù", "bù khí", "cấp gió bù", "bù gió"))
                and not any(k in raw_makeup_chk for k in ("fag", "sag", "rag", "oag", "pag", "seg", "eag", "sdg", "obd", "nan thẳng", "bầu dục"))
            )
            is_ktc = (bool(re.search(ktc_rx, dim_search_text, re.IGNORECASE)) or is_makeup_air_z) and not has_explicit_kt_bi

            # Rule 6 & 7: Kiểm tra dòng sản phẩm kết hợp Van + Cửa (Combo)
            search_combo_txt = re.sub(r'\(\s*(?:for|cho)\s+[^)]*\)', '', f"{raw_name_str} {name_str} {model_str}".lower().replace("\n", " "))
            diff_rx = combo_cfg.get("combo_diffuser_regex", r'\bmiệng\s*gió\b|\bcửa\s*gió\b|\bcửa\s*nan\b|\blouver\s*|\bluover\b|\blouvre\b|\bgrille\b|\bgriller\b|\bdiffuser\b|cửa\s*(?:chớp|nan|thẳng|gió|hút|cấp|hồi|thải|lấy|xả|tăng\s*áp|tạo\s*áp|\+)')
            damp_rx = combo_cfg.get("combo_damper_regex", r'\b(mfd|mfsd|md|fd|prd|nrd|vcd)\b|van\s+(?:on/off|on-off|chống\s*cháy|chặn\s*lửa|ngăn\s*cháy|động\s*cơ|một\s*chiều|1\s*chiều|xả\s*áp)')
            is_access_door_item = any(k in search_combo_txt for k in ("access door", "cửa công tác", "cửa thăm", "cửa kiểm tra"))
            has_diffuser = bool(re.search(diff_rx, search_combo_txt)) and not is_access_door_item
            m_damper = re.search(damp_rx, search_combo_txt) if not is_access_door_item else None

            damper_priority = {"MFSD": 1, "MFD": 2, "FD": 3, "MD": 4, "PRD": 5, "NRD": 6, "VCD": 7}

            def _detect_damper_type(tok_str):
                t = tok_str.strip().lower()
                if re.search(diff_rx, t) and not re.search(r'\b(mfsd|mfd|vcd|vd|nrd|prd|md|fd)\b', t):
                    m_v = re.search(r'\bvan\b.*', t)
                    if not m_v:
                        return None
                    t = m_v.group(0)
                if re.search(r'\bmfsd\b', t): return "MFSD"
                if re.search(r'\bmfd\b|ngăn\s*cháy.*động\s*cơ|chống\s*cháy.*động\s*cơ|chặn\s*lửa.*động\s*cơ', t): return "MFD"
                if re.search(r'\bvcd\b|\bvd\b|điều\s*chỉnh.*lưu\s*lượng', t): return "VCD"
                if re.search(r'\bnrd\b|1\s*chiều|một\s*chiều', t): return "NRD"
                if re.search(r'\bprd\b|xả\s*áp|giảm\s*áp', t): return "PRD"
                if re.search(r'\bmd\b|động\s*cơ|on/off|on-off', t): return "MD"
                if re.search(r'\bfd\b|chặn\s*lửa|chống\s*cháy|dập\s*lửa|cầu\s*chì', t): return "FD"
                return None

            scan_base = raw_name_str if re.search(r'[a-zA-ZÀ-ỹ]{2,}', re.sub(r'\b(?:mm|kt|w|h|l|d|dn|phi)\b', '', raw_name_str, flags=re.IGNORECASE)) else name_str
            scan_source = f"{scan_base} + {model_str}" if model_str else scan_base
            tokens = [tk.strip() for tk in re.split(r'\s*(?:\+|(?:\bkèm\b|\bgắn\b|\btích\s*hợp\b|\bkết\s*hợp\b|\bc/w\.?\b|\bwith\b)(?!\s*(?:động\s*cơ|mô\s*tơ|môtơ|mo\s*tơ|motor|cầu\s*chì)))\s*', scan_source, flags=re.IGNORECASE) if tk.strip()]
            detected_dampers = []
            for tk in tokens:
                dt = _detect_damper_type(tk)
                if dt and dt not in detected_dampers:
                    detected_dampers.append(dt)
            if not detected_dampers and m_damper:
                dt = _detect_damper_type(m_damper.group(0))
                if dt:
                    detected_dampers.append(dt)

            if category in ("CUA_GIO", "VAN_GIO") and has_diffuser and m_damper and detected_dampers:
                sorted_dampers = sorted(detected_dampers, key=lambda x: damper_priority.get(x, 99))
                w_diff = (w + 50.0) if (is_ktc and w > 0) else w
                h_diff = (h + 50.0) if (is_ktc and h > 0) else h
                d_diff = (d_val + 50.0) if (is_ktc and d_val > 0) else d_val
                items.append({
                    "source_row": r,
                    "source_sheet": sheet_name,
                    "section": current_section,
                    "raw_name": raw_name_str,
                    "name": name_str,
                    "aux_text": aux_context,
                    "note_boq": note_boq,
                    "category": "CUA_GIO",
                    "combo_damper": sorted_dampers[0],
                    "combo_dampers": sorted_dampers,
                    "w_damper": w,
                    "h_damper": h,
                    "d_damper": d_val,
                    "model": model_str.upper(),
                    "w": w_diff,
                    "h": h_diff,
                    "d": d_diff,
                    "l": l,
                    "thickness": 0.0,
                    "ei_rating": ei_rating,
                    "unit": str(unit_val).strip() if str(unit_val).strip() else "cái",
                    "qty": float(qty_val),
                    "is_missing_qty": is_missing_qty,
                    "ref_price": ws.cell(r, col_map["price"]).value if "price" in col_map else None,
                    "is_ktc_applied": bool(is_ktc),
                    "has_embedded_image": (r in image_rows)
                })

            elif category == "VAN_GIO" and not has_diffuser:
                if len(detected_dampers) > 1:
                    # Sắp xếp theo thứ bậc ưu tiên để chọn van chính (MFSD > MFD > FD > MD > PRD > NRD > VCD)
                    sorted_dampers = sorted(detected_dampers, key=lambda x: damper_priority.get(x, 99))
                    primary_type = sorted_dampers[0]
                    secondary_dampers = sorted_dampers[1:]
                    secondary_type = secondary_dampers[0] if secondary_dampers else None

                    items.append({
                        "source_row": r,
                        "source_sheet": sheet_name,
                        "section": current_section,
                        "raw_name": raw_name_str,
                        "name": name_str,
                        "aux_text": aux_context,
                        "note_boq": note_boq,
                        "category": "VAN_GIO",
                        "model": primary_type,
                        "secondary_damper": secondary_type,
                        "secondary_dampers": secondary_dampers,
                        "w": w,
                        "h": h,
                        "d": d_val,
                        "l": l,
                        "thickness": thickness,
                        "ei_rating": ei_rating,
                        "unit": str(unit_val).strip() if str(unit_val).strip() else "cái",
                        "qty": float(qty_val),
                        "is_missing_qty": is_missing_qty,
                        "ref_price": ws.cell(r, col_map["price"]).value if "price" in col_map else None
                    })
                else:
                    # Không match combo Van+Van → xử lý như van đơn thông thường
                    item_data = {
                        "source_row": r,
                        "source_sheet": sheet_name,
                        "section": current_section,
                        "raw_name": raw_name_str,
                        "name": name_str,
                        "aux_text": aux_context,
                        "note_boq": note_boq,
                        "category": category,
                        "orig_category": orig_category,
                        "is_missing_dims": is_missing_dims,
                        "model": model_str.upper(),
                        "w": w,
                        "h": h,
                        "w2": w2,
                        "h2": h2,
                        "d": d_val,
                        "l": l,
                        "thickness": thickness,
                        "flange_type": fitting_geom.get("flange_type") if fitting_geom else None,
                        "ei_rating": ei_rating,
                        "unit": str(unit_val).strip(),
                        "qty": float(qty_val),
                        "is_missing_qty": is_missing_qty,
                        "ref_price": ws.cell(r, col_map["price"]).value if "price" in col_map else None,
                        "is_ktc_applied": False,
                        "fitting_geom": fitting_geom
                    }
                    items.append(item_data)

            else:
                non_quotable_cats = ("THUONG_MAI", "KHAC") if self.include_flex_duct else ("THUONG_MAI", "ONG_MEM", "KHAC")
                if category in non_quotable_cats or (w == 0 and h == 0 and d_val == 0 and category in ("ONG_GIO", "TIEU_AM", "CANVAS")):
                    ref_p = ws.cell(r, col_map["price"]).value if "price" in col_map else None
                    if (not isinstance(ref_p, (int, float)) or ref_p <= 0) and not (not name_lower.startswith("quạt ") and any(k in name_lower for k in ("chống cháy ei", "lưới bảo vệ"))):
                        self.skipped_rows[sheet_name].add(r)
                        raw_spec_cell = ws.cell(r, col_map["spec"]).value if "spec" in col_map else None
                        raw_spec_disp = raw_spec_cell if raw_spec_cell not in (None, "") else (spec_str or None)
                        self.skip_reasons.setdefault(sheet_name, {})[r] = {
                            "reason": "MISSING_DIMS" if is_missing_dims else (f"NO_DIMS_{category}" if category in ("ONG_GIO", "TIEU_AM", "CANVAS") else f"OUT_OF_SCOPE_{category}"),
                            "category": category,
                            "orig_category": orig_category,
                            "name": raw_name_str,
                            "spec": raw_spec_disp,
                            "qty": _safe_qty(qty_val),
                            "unit": unit_val,
                        }
                        if is_missing_dims and float(qty_val or 0) > 0:
                            self.missing_dim_rows[sheet_name].append(f"Dòng {r}: Tên='{raw_name_str}' | Spec='{raw_spec_disp}' | KL={float(qty_val)}")
                        continue

                # Không phải Combo: Xử lý KTC cho Cửa gió (độ dày & chiều dài Van gió lấy trực tiếp từ Agent05_DamperBOM)
                if category == "CUA_GIO" and is_ktc:
                    if w_neck == 0 and w > 0: w_neck = w
                    if h_neck == 0 and h > 0: h_neck = h
                    if w > 0: w += 50.0
                    if h > 0: h += 50.0
                    if d_val > 0: d_val += 50.0

                item_data = {
                    "source_row": r,
                    "source_sheet": sheet_name,
                    "section": current_section,
                    "raw_name": raw_name_str,
                    "name": name_str,
                    "aux_text": aux_context,
                    "note_boq": note_boq,
                    "category": category,
                    "orig_category": orig_category,
                    "is_missing_dims": is_missing_dims,
                    "model": model_str.upper(),
                    "w": w,
                    "h": h,
                    "w2": w2,
                    "h2": h2,
                    "w_neck": w_neck,
                    "h_neck": h_neck,
                    "d": d_val,
                    "l": l,
                    "thickness": thickness,
                    "flange_type": fitting_geom.get("flange_type") if fitting_geom else None,
                    "ei_rating": ei_rating,
                    "unit": str(unit_val).strip(),
                    "qty": float(qty_val),
                    "is_missing_qty": is_missing_qty,
                    "ref_price": ws.cell(r, col_map["price"]).value if "price" in col_map else None,
                    "is_ktc_applied": bool(category == "CUA_GIO" and is_ktc),
                    "has_embedded_image": (r in image_rows),
                    "fitting_geom": fitting_geom
                }
                items.append(item_data)
            
            if max_items and len(items) >= max_items:
                break

        # Quét bảng phụ song song bên phải (Side-by-Side Table, VD: VAN GIÓ TƯƠI bên trái + VAN HÚT KHÓI bên phải)
        if not _is_side_table and not getattr(ws, "_forced_qty_col", None):
            from openpyxl.utils import get_column_letter
            max_main_col = max((v for k, v in col_map.items() if isinstance(v, int)), default=3)
            used_side_cols = set()
            for c_scan in range(max_main_col + 1, max_scan_col + 1):
                if c_scan in used_side_cols or get_column_letter(c_scan) in hidden_cols:
                    continue
                matched_rows = []
                for rr in range(1, min(60, ws.max_row + 1)):
                    if rr in hidden_rows:
                        continue
                    v_nm = ws.cell(rr, c_scan).value
                    if not isinstance(v_nm, str):
                        continue
                    v_s = unicodedata.normalize('NFC', v_nm).strip()
                    if re.search(r'(\d{2,4}\s*[xX×*]\s*\d{2,4}|\b[dDØ]\d{2,4}\b)', v_s) and re.search(r'\b(van|mfd|mfsd|fd|vcd|md|nrd|prd|obd|cửa|miệng|louver|ống|cút|côn|tê)\b', v_s, re.IGNORECASE):
                        for c_q in range(c_scan + 1, min(c_scan + 4, max_scan_col + 1)):
                            vq = ws.cell(rr, c_q).value
                            if isinstance(vq, (int, float)) and vq > 0:
                                matched_rows.append((rr, c_q))
                                break
                if len(matched_rows) >= 2:
                    first_side_r = matched_rows[0][0]
                    c_qty_side = matched_rows[0][1]
                    side_hdr_r = max(1, first_side_r - 1)
                    c_stt_side = c_scan - 1 if (c_scan - 1 > max_main_col and ws.cell(first_side_r, c_scan - 1).value is not None) else 0
                    side_col_map = {"name": c_scan, "qty": c_qty_side}
                    if c_stt_side > 0:
                        side_col_map["stt"] = c_stt_side
                    if c_qty_side > c_scan + 1:
                        side_col_map["unit"] = c_scan + 1

                    def _find_table_title(c_from, c_to, max_r):
                        for tr in range(1, max_r + 1):
                            for tc in range(c_from, c_to + 1):
                                tv = ws.cell(tr, tc).value
                                if isinstance(tv, str):
                                    ts = unicodedata.normalize('NFC', tv).strip()
                                    tl = ts.lower()
                                    if len(ts) > 2 and not re.match(r'^(?:\(\d+\)|cột\s*\d+|column\s*\d+|stt|khu\s*vực|vị\s*trí|sl|số\s*lượng|kích\s*thước.*|đvt|đơn\s*vị)$', tl):
                                        return ts
                        return None

                    main_title = _find_table_title(1, max_main_col, max(1, header_row - 1))
                    if main_title:
                        for it_m in items:
                            if not it_m.get("is_side_table") and it_m.get("section") in (None, "", "SẢN PHẨM"):
                                it_m["section"] = main_title
                                it_m["name"] = f"{main_title} {it_m['raw_name']}"
                    side_title = _find_table_title(max(max_main_col + 1, c_stt_side or c_scan), c_qty_side, side_hdr_r) or "HẠNG MỤC BỔ SUNG"
                    side_items = self.parse_sheet(
                        sheet_name,
                        max_items=max_items,
                        _override_col_map=side_col_map,
                        _override_header_row=side_hdr_r,
                        _initial_section=side_title,
                        _is_side_table=True
                    )
                    for s_it in side_items:
                        s_it["is_side_table"] = True
                        s_it["side_stt"] = ws.cell(s_it["source_row"], c_stt_side).value if c_stt_side > 0 else None
                    items.extend(side_items)
                    for cc in range(c_stt_side or c_scan, c_qty_side + 2):
                        used_side_cols.add(cc)
                
        if not _is_side_table and _override_col_map is None and _override_header_row is None and max_items is None and not self.row_overrides:
            if not hasattr(self.wb, "_swind_parsed_cache"):
                self.wb._swind_parsed_cache = {}
            self.wb._swind_parsed_cache[(sheet_name, self.include_flex_duct)] = (
                [dict(it) for it in items],
                set(self.skipped_rows.get(sheet_name, set())),
                list(self.missing_dim_rows.get(sheet_name, [])),
                dict(self.skip_reasons.get(sheet_name, {}))
            )
        return items

    def auto_detect_hvac_sheets(self):
        """
        Tự động nhận diện các Sheet liên quan đến Cơ Điện Lạnh (HVAC/MEP).
        """
        visible_sheets = [s for s in self.wb.sheetnames if getattr(self.wb[s], "sheet_state", "visible") == "visible"]
        cand_sheets = visible_sheets if visible_sheets else list(self.wb.sheetnames)
        zone_sheets = [s for s in cand_sheets if getattr(self.wb[s], "_forced_qty_col", None)]
        if zone_sheets:
            return zone_sheets

        combo_cfg = get_subtype_rules().get("COMBO_AND_DETECTION_RULES", {})
        keywords = combo_cfg.get("hvac_sheet_keywords", [
            "dhkk", "hvac", "thong gio", "thông gió", "thong_gio", 
            "hut khoi", "hút khói", "hut_khoi", "tang ap", "tăng áp", "tang áp", "tăng ap",
            "tang_ap", "duct", "lanh", "lạnh", "ventilation", "m.4 duct"
        ])
        detected = []
        for s in cand_sheets:
            s_lower = s.lower().strip()
            if any(kw in s_lower for kw in keywords):
                detected.append(s)
        return detected

    def parse_sheets(self, sheet_names=None, max_items_per_sheet=None):
        """
        Quét nhiều sheet dự thầu và gom cụm sản phẩm.
        """
        if sheet_names is None:
            sheet_names = self.auto_detect_hvac_sheets()
            if not sheet_names:
                visible_sheets = [s for s in self.wb.sheetnames if getattr(self.wb[s], "sheet_state", "visible") == "visible"]
                cand_sheets = visible_sheets if visible_sheets else list(self.wb.sheetnames)
                # Nếu không auto-detect được sheet nào, lấy sheet đầu tiên không phải summary
                for s in cand_sheets:
                    if not any(k in s.lower() for k in ("summary", "cover", "bìa", "phạm vi", "pham vi", "scope")):
                        sheet_names = [s]
                        break
                if not sheet_names:
                    sheet_names = [cand_sheets[0]]

        all_items = []
        # Parse tuần tự: các sheet dùng chung 1 workbook openpyxl (ws.cell() có thể tạo ô mới -> race
        # "dictionary changed size during iteration" làm mất trắng cả sheet khi chạy đa luồng).
        for s in sheet_names:
            try:
                sheet_items = self.parse_sheet(s, max_items=max_items_per_sheet)
                all_items.extend(sheet_items)
            except Exception as e:
                self.parse_errors[s] = f"{type(e).__name__}: {e}"
                print(f"[SwindBOQParser] ⛔ LỖI bóc tách sheet '{s}' (sheet bị bỏ qua — báo giá CHƯA HOÀN CHỈNH): {e}")
        return all_items
