import xlrd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import json
import re
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_FILE = os.path.join(BASE_DIR, "BẢNG TÍNH GIÁ SWIND  chuẩn 2026.xls")
MASTER_DIR = os.path.join(BASE_DIR, "data", "master_data")
OUTPUT_EXCEL_TECH = os.path.join(MASTER_DIR, "MASTER_DATA_TECH.xlsx")
OUTPUT_EXCEL_PRICE = os.path.join(MASTER_DIR, "MASTER_DATA_PRICE.xlsx")
OUTPUT_JSON_TECH = os.path.join(MASTER_DIR, "master_database_tech.json")
OUTPUT_JSON_PRICE = os.path.join(MASTER_DIR, "master_database_price.json")

def to_ascii_slug(text):
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.replace('đ', 'd').replace('Đ', 'D')
    text = re.sub(r'<\s*300', '-LT-300', text)
    text = re.sub(r'>\s*300', '-GT-300', text)
    text = re.sub(r'[^a-zA-Z0-9]+', '-', text).strip('-').upper()
    return text

def clean_str(val):
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return str(val)
    return str(val).strip()

def clean_num(val):
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        val = val.replace(',', '').strip()
        try:
            return float(val)
        except ValueError:
            return None
    return None

def extract_all_master_data():
    wb = xlrd.open_workbook(INPUT_FILE, formatting_info=False)
    master_records = []
    
    # -------------------------------------------------------------
    # 1. SHEET: Vat tu tôn
    # -------------------------------------------------------------
    s_ton = wb.sheet_by_name('Vat tu tôn')
    
    # Quét nhận diện các Block ở Row 1
    blocks = {}
    for c in range(s_ton.ncols):
        val = clean_str(s_ton.cell_value(1, c))
        if val:
            blocks[c] = val
            
    col_keys = sorted(blocks.keys())
    
    for i, start_col in enumerate(col_keys):
        end_col = col_keys[i+1] if i+1 < len(col_keys) else s_ton.ncols
        block_name = blocks[start_col]
        
        # Đọc subheaders ở Row 2
        subheaders = {c: clean_str(s_ton.cell_value(2, c)).lower() for c in range(start_col, end_col)}
        
        # Phân nhánh theo loại block
        if "TÔN" in block_name.upper():
            brand_code = "DONGBU" if "DONGBU" in block_name else \
                         "VNSTEEL" if "VNSTEEL" in block_name else \
                         "PN" if "PHƯƠNG NAM" in block_name else \
                         "HOASEN" if "HOA SEN" in block_name else \
                         "BLUESCOPE" if "BLUESCOPE" in block_name else "OTHER"
                         
            # Tìm cột độ dày, giá/m2, giá/kg, tỉ trọng
            c_thk = next((c for c, h in subheaders.items() if "dày" in h), None)
            c_p_m2 = next((c for c, h in subheaders.items() if "m2" in h and "đơn giá" in h), None)
            c_p_kg = next((c for c, h in subheaders.items() if "kg" in h and "đơn giá" in h), None)
            c_dens = next((c for c, h in subheaders.items() if "trọng" in h), None)
            
            for r in range(4, s_ton.nrows):
                if c_thk is None: break
                thk = clean_num(s_ton.cell_value(r, c_thk))
                p_m2 = clean_num(s_ton.cell_value(r, c_p_m2)) if c_p_m2 is not None else None
                p_kg = clean_num(s_ton.cell_value(r, c_p_kg)) if c_p_kg is not None else None
                dens = clean_num(s_ton.cell_value(r, c_dens)) if c_dens is not None else None
                
                # Barem tôn: nếu không có cột tỷ trọng thì tính theo 7.85 * thickness
                if thk and thk > 0 and (p_m2 or p_kg):
                    calc_dens = dens if dens and dens > 0 else round(7.85 * thk, 3)
                    # Nếu có giá m2 mà chưa có giá kg -> giá kg = giá m2 / dens
                    if p_m2 and not p_kg and calc_dens > 0:
                        p_kg = round(p_m2 / calc_dens, 2)
                    elif p_kg and not p_m2 and calc_dens > 0:
                        p_m2 = round(p_kg * calc_dens, 2)
                        
                    mat_id = f"MAT-STEEL-{brand_code}-{thk:.2f}"
                    master_records.append({
                        "MAT_ID": mat_id,
                        "Name": f"{block_name} {thk:.2f}mm",
                        "Category": "TÔN_CUỘN_MẠ_KẼM",
                        "Thickness_mm": thk,
                        "Unit": "kg",
                        "Price": p_kg,
                        "Price_m2": p_m2,
                        "Density_kg_m2": calc_dens,
                        "Source_Trace": f"Vat tu tôn!R{r+1}C{c_thk+1}"
                    })
                    if brand_code == "PN" and abs(thk - 1.15) < 0.01:
                        master_records.append({
                            "MAT_ID": "MAT-STEEL-PN-1.20",
                            "Name": f"{block_name} 1.20mm",
                            "Category": "TÔN_CUỘN_MẠ_KẼM",
                            "Thickness_mm": 1.20,
                            "Unit": "kg",
                            "Price": p_kg,
                            "Price_m2": p_m2,
                            "Density_kg_m2": calc_dens,
                            "Source_Trace": f"Vat tu tôn!R{r+1}C{c_thk+1}"
                        })
                    elif brand_code == "PN" and abs(thk - 1.48) < 0.01:
                        master_records.append({
                            "MAT_ID": "MAT-STEEL-PN-1.50",
                            "Name": f"{block_name} 1.50mm",
                            "Category": "TÔN_CUỘN_MẠ_KẼM",
                            "Thickness_mm": 1.50,
                            "Unit": "kg",
                            "Price": p_kg,
                            "Price_m2": p_m2,
                            "Density_kg_m2": calc_dens,
                            "Source_Trace": f"Vat tu tôn!R{r+1}C{c_thk+1}"
                        })
                    elif brand_code == "HOASEN" and abs(thk - 1.15) < 0.01:
                        master_records.append({
                            "MAT_ID": "MAT-STEEL-HOA_SEN-1.15",
                            "Name": f"{block_name} 1.15mm",
                            "Category": "TÔN_CUỘN_MẠ_KẼM",
                            "Thickness_mm": 1.15,
                            "Unit": "kg",
                            "Price": p_kg,
                            "Price_m2": p_m2,
                            "Density_kg_m2": calc_dens,
                            "Source_Trace": f"Vat tu tôn!R{r+1}C{c_thk+1}"
                        })
                    
        elif "INOX" in block_name.upper():
            c_name = next((c for c, h in subheaders.items() if "inox" in h or "loại" in h), None)
            c_p_kg = next((c for c, h in subheaders.items() if "kg" in h and "đơn giá" in h), None)
            c_p_m2 = next((c for c, h in subheaders.items() if "m2" in h and "đơn giá" in h), None)
            
            for r in range(4, s_ton.nrows):
                if c_name is None: break
                name = clean_str(s_ton.cell_value(r, c_name))
                p_kg = clean_num(s_ton.cell_value(r, c_p_kg)) if c_p_kg is not None else None
                p_m2 = clean_num(s_ton.cell_value(r, c_p_m2)) if c_p_m2 is not None else None
                
                if name and (p_kg or p_m2):
                    slug = to_ascii_slug(name)
                    master_records.append({
                        "MAT_ID": f"MAT-INOX-{slug}",
                        "Name": f"Inox {name}",
                        "Category": "INOX",
                        "Unit": "kg" if p_kg else "m2",
                        "Price": p_kg if p_kg else p_m2,
                        "Price_m2": p_m2,
                        "Source_Trace": f"Vat tu tôn!R{r+1}C{c_name+1}"
                    })
                    
        elif "PHỤ VAN" in block_name.upper():
            c_name = next((c for c, h in subheaders.items() if "vật tư" in h or "loại" in h), None)
            c_unit = next((c for c, h in subheaders.items() if "đơn vị" in h), None)
            c_price = next((c for c, h in subheaders.items() if "đơn giá" in h), None)
            
            for r in range(4, s_ton.nrows):
                if c_name is None: break
                name = clean_str(s_ton.cell_value(r, c_name))
                unit = clean_str(s_ton.cell_value(r, c_unit)) if c_unit is not None else "Cái"
                price = clean_num(s_ton.cell_value(r, c_price)) if c_price is not None else None
                
                if name and price and price > 0:
                    slug = to_ascii_slug(name)
                    master_records.append({
                        "MAT_ID": f"MAT-DAMPER-{slug}",
                        "Name": name,
                        "Category": "PHỤ_KIỆN_VAN",
                        "Unit": unit if unit else "Cái",
                        "Price": price,
                        "Source_Trace": f"Vat tu tôn!R{r+1}C{c_name+1}"
                    })

        elif "KE + BÍCH" in block_name.upper():
            c_name = next((c for c, h in subheaders.items() if "vật tư" in h or "loại" in h), None)
            c_unit = next((c for c, h in subheaders.items() if "đơn vị" in h), None)
            c_price = next((c for c, h in subheaders.items() if "đơn giá" in h), None)
            
            for r in range(4, s_ton.nrows):
                if c_name is None: break
                name = clean_str(s_ton.cell_value(r, c_name))
                unit = clean_str(s_ton.cell_value(r, c_unit)) if c_unit is not None else "m"
                price = clean_num(s_ton.cell_value(r, c_price)) if c_price is not None else None
                
                if name and price and price > 0:
                    slug = to_ascii_slug(name)
                    master_records.append({
                        "MAT_ID": f"MAT-FLANGE-{slug}",
                        "Name": name,
                        "Category": "THÉP_BÍCH_GÓC",
                        "Unit": unit if unit else "m",
                        "Price": price,
                        "Source_Trace": f"Vat tu tôn!R{r+1}C{c_name+1}"
                    })

        elif "BẢO ÔN + TIÊU ÂM" in block_name.upper():
            c_name = next((c for c, h in subheaders.items() if "vật tư" in h or "loại" in h), None)
            c_unit = next((c for c, h in subheaders.items() if "đơn vị" in h), None)
            c_price = next((c for c, h in subheaders.items() if "đơn giá" in h), None)
            
            for r in range(4, s_ton.nrows):
                if c_name is None: break
                name = clean_str(s_ton.cell_value(r, c_name))
                unit = clean_str(s_ton.cell_value(r, c_unit)) if c_unit is not None else "m2"
                price = clean_num(s_ton.cell_value(r, c_price)) if c_price is not None else None
                
                if name and price and price > 0:
                    slug = to_ascii_slug(name)
                    master_records.append({
                        "MAT_ID": f"MAT-INSUL-{slug}",
                        "Name": name,
                        "Category": "VẬT_LIỆU_TIÊU_ÂM",
                        "Unit": unit if unit else "m2",
                        "Price": price,
                        "Source_Trace": f"Vat tu tôn!R{r+1}C{c_name+1}"
                    })

    # -------------------------------------------------------------
    # 2. SHEET: BG Cửa  (rút gọn) - Profile nhôm & phụ kiện cửa
    # -------------------------------------------------------------
    s_cua = wb.sheet_by_name('BG Cửa  (rút gọn)')
    for r in range(2, 19):
        name = clean_str(s_cua.cell_value(r, 1))
        weight = clean_num(s_cua.cell_value(r, 4))
        price_old = clean_num(s_cua.cell_value(r, 5))
        price_new = clean_num(s_cua.cell_value(r, 6))
        code = clean_str(s_cua.cell_value(r, 7))
        price = price_new if price_new else price_old
        
        if name and weight and price:
            slug = to_ascii_slug(name)
            if code:
                slug += f"-{code.replace(' ', '')}"
            master_records.append({
                "MAT_ID": f"MAT-ALU-{slug}",
                "Name": f"Nhôm định hình {name} {code}".strip(),
                "Category": "NHÔM_ĐỊNH_HÌNH",
                "Unit": "kg",
                "Price": price,
                "Weight_kg_m": weight,
                "Source_Trace": f"BG Cửa  (rút gọn)!R{r+1}C2"
            })
            
    # Sơn tĩnh điện & Phụ kiện từ BG Cửa (rút gọn)
    # R31C20: Sơn tĩnh điện 75,000 đ/m2
    # R33C19: Phụ kiện cửa gió 15,000 đ/bộ
    master_records.append({
        "MAT_ID": "MAT-PAINT-POWDER",
        "Name": "Sơn tĩnh điện tiêu chuẩn Polyester",
        "Category": "SƠN_TĨNH_ĐIỆN",
        "Unit": "m2",
        "Price": 75000.0,
        "Source_Trace": "BG Cửa  (rút gọn)!R32C21"
    })
    master_records.append({
        "MAT_ID": "MAT-DOOR-ACCESSORY",
        "Name": "Phụ kiện lắp ghép / Ke góc cửa gió",
        "Category": "PHỤ_KIỆN_CỬA",
        "Unit": "Bộ",
        "Price": 15000.0,
        "Source_Trace": "BG Cửa  (rút gọn)!R34C20"
    })

    # -------------------------------------------------------------
    # 3. SHEET: VT van - Động cơ & Cầu chì chi tiết
    # -------------------------------------------------------------
    s_van = wb.sheet_by_name('VT van')
    header_r = None
    for r in range(min(20, s_van.nrows)):
        row_str = " ".join([clean_str(s_van.cell_value(r, c)).lower() for c in range(s_van.ncols)])
        if "sản phẩm" in row_str and "đơn giá" in row_str:
            header_r = r
            break
            
    if header_r is not None:
        headers_van = {c: clean_str(s_van.cell_value(header_r, c)).lower() for c in range(s_van.ncols)}
        c_name = next(c for c, h in headers_van.items() if "sản phẩm" in h)
        c_price = next(c for c, h in headers_van.items() if "đơn giá" in h)
        c_brand = next((c for c, h in headers_van.items() if "nhãn" in h or "xuất xứ" in h), None)
        
        for r in range(header_r + 1, s_van.nrows):
            name = clean_str(s_van.cell_value(r, c_name))
            price = clean_num(s_van.cell_value(r, c_price))
            brand = clean_str(s_van.cell_value(r, c_brand)) if c_brand else ""
            
            if name and price and price > 0:
                model_match = re.search(r'(NACA\s*[\d\.\-]+|NACM\s*[\d\.\-]+|NTC\s*\d+|SAFA\s*[\d\s\-S]+|NPG\s*\d+|Model:\s*[AB])', name, re.IGNORECASE)
                if model_match:
                    model_str = to_ascii_slug(model_match.group(1))
                else:
                    model_str = to_ascii_slug(name)[:20]
                    
                mat_id = f"MAT-ACTUATOR-{model_str}"
                master_records.append({
                    "MAT_ID": mat_id,
                    "Name": name.replace('\n', ' ').strip(),
                    "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN",
                    "Unit": "Cái",
                    "Price": price,
                    "Brand": brand if brand else "NENUTEC",
                    "Source_Trace": f"VT van!R{r+1}C{c_name+1}"
                })

    # -------------------------------------------------------------
    # 4. SHEET: VT tiêu âm - Bông thủy tinh, bông khoáng, tôn lỗ
    # -------------------------------------------------------------
    s_ta = wb.sheet_by_name('VT tiêu âm')
    header_r_ta = None
    for r in range(min(15, s_ta.nrows)):
        row_str = " ".join([clean_str(s_ta.cell_value(r, c)).lower() for c in range(s_ta.ncols)])
        if "sản phẩm" in row_str and "đơn giá" in row_str:
            header_r_ta = r
            break
            
    if header_r_ta is not None:
        headers_ta = {c: clean_str(s_ta.cell_value(header_r_ta, c)).lower() for c in range(s_ta.ncols)}
        c_name = next(c for c, h in headers_ta.items() if "sản phẩm" in h)
        c_price = next(c for c, h in headers_ta.items() if "đơn giá" in h)
        c_unit = next((c for c, h in headers_ta.items() if "đơn vị" in h), None)
        
        for r in range(header_r_ta + 1, s_ta.nrows):
            name = clean_str(s_ta.cell_value(r, c_name))
            unit = clean_str(s_ta.cell_value(r, c_unit)) if c_unit is not None else "m2"
            price = clean_num(s_ta.cell_value(r, c_price))
            
            if name and price and price > 0:
                slug = to_ascii_slug(name.split('\n')[0])[:25].strip('-')
                master_records.append({
                    "MAT_ID": f"MAT-INSUL-{slug}-{r+1}",
                    "Name": name.replace('\n', ' - ').strip(),
                    "Category": "VẬT_LIỆU_TIÊU_ÂM",
                    "Unit": unit if unit else "m2",
                    "Price": price,
                    "Source_Trace": f"VT tiêu âm!R{r+1}C{c_name+1}"
                })

    # -------------------------------------------------------------
    # 5. SHEET: BG canvas, ventcap & BG TIEU AM & bảo ôn & BG VAN CU
    # -------------------------------------------------------------
    if 'BG canvas, ventcap' in wb.sheet_names():
        s_cv = wb.sheet_by_name('BG canvas, ventcap')
        cv_price = clean_num(s_cv.cell_value(5, 13)) or 57600.0
        master_records.append({
            "MAT_ID": "CANVAS_FIRE_RESISTANT",
            "Name": "Cuộn vải bạt bọc cách nhiệt/chống cháy",
            "Category": "VẬT_TƯ_BẠT",
            "Unit": "m2",
            "Price": cv_price,
            "Source_Trace": "BG canvas, ventcap"
        })

    if 'BG TIEU AM' in wb.sheet_names():
        s_bgta = wb.sheet_by_name('BG TIEU AM')
        perf_price = clean_num(s_bgta.cell_value(15, 20)) or 55000.0
        rw_price = clean_num(s_bgta.cell_value(15, 21)) or 65000.0
        gc_price = clean_num(s_bgta.cell_value(15, 22)) or 15000.0
        master_records.extend([
            {"MAT_ID": "ROCKWOOL_50MM_50KG", "Category": "VẬT_LIỆU_TIÊU_ÂM", "Name": "Bông Rockwool tỷ trọng 50kg/m3 dày 50mm", "Unit": "m2", "Price": rw_price, "Source_Trace": "BG TIEU AM"},
            {"MAT_ID": "GLASS_CLOTH", "Category": "VẬT_LIỆU_TIÊU_ÂM", "Name": "Vải thủy tinh bọc tiêu âm", "Unit": "m2", "Price": gc_price, "Source_Trace": "BG TIEU AM"},
            {"MAT_ID": "PERFORATED_SHEET", "Category": "VẬT_LIỆU_TIÊU_ÂM", "Name": "Tôn đục lỗ tiêu âm", "Unit": "m2", "Price": perf_price, "Source_Trace": "BG TIEU AM"}
        ])

    if 'bảo ôn' in wb.sheet_names():
        s_bo = wb.sheet_by_name('bảo ôn')
        master_records.append({
            "MAT_ID": "SUPERLON_19MM",
            "Category": "VẬT_LIỆU_BẢO_ÔN",
            "Name": "Bảo ôn Superlon dạng tấm dày 19mm",
            "Unit": "m2",
            "Price": 220000.0,
            "Source_Trace": "bảo ôn"
        })
        for r in range(14, min(26, s_bo.nrows)):
            item_name = clean_str(s_bo.cell_value(r, 1))
            size_str = clean_str(s_bo.cell_value(r, 3))
            unit = clean_str(s_bo.cell_value(r, 9)) if s_bo.ncols > 9 else "ống"
            price = clean_num(s_bo.cell_value(r, 11))
            if item_name and size_str and price and price > 0:
                clean_id = f"FLEX_DUCT_{size_str}_{'INSUL' if 'có' in item_name else 'NON_INSUL'}"
                master_records.append({
                    "MAT_ID": clean_id,
                    "Category": "ỐNG_GIÓ_MỀM",
                    "Name": item_name,
                    "Unit": unit if unit else "ống",
                    "Price": price,
                    "Source_Trace": "bảo ôn"
                })
        master_records.extend([
            {"MAT_ID": "FLEX_DUCT_D400_NON_INSUL", "Category": "ỐNG_GIÓ_MỀM", "Name": "ống mềm D400 không bảo ôn", "Unit": "ống", "Price": 225000.0, "Source_Trace": "bảo ôn"},
            {"MAT_ID": "FLEX_DUCT_D400_INSUL", "Category": "ỐNG_GIÓ_MỀM", "Name": "ống mềm D400 có bảo ôn", "Unit": "ống", "Price": 320000.0, "Source_Trace": "bảo ôn"}
        ])

    # 6. SHEET: BG VAN  - Tấm chống cháy EI, Động cơ Belimo/Honeywell/Swind & Phụ kiện FD/MFD
    if 'BG VAN ' in wb.sheet_names():
        s_bgvan = wb.sheet_by_name('BG VAN ')
        p_fslf = clean_num(s_bgvan.cell_value(15, 40)) or 3700000.0
        p_fslf_buy = clean_num(s_bgvan.cell_value(11, 40)) or 3315700.0
        p_bfn = clean_num(s_bgvan.cell_value(15, 41)) or 3700000.0
        p_bfn_buy = clean_num(s_bgvan.cell_value(11, 41)) or 2973700.0
        p_fb10_body = clean_num(s_bgvan.cell_value(15, 42)) or 700000.0
        p_fb20_blade = clean_num(s_bgvan.cell_value(15, 43)) or 800000.0
        p_fb10_blade = clean_num(s_bgvan.cell_value(15, 44)) or 700000.0
        p_honeywell = clean_num(s_bgvan.cell_value(15, 45)) or 3600000.0
        p_honeywell_buy = clean_num(s_bgvan.cell_value(11, 45)) or 2850000.0
        p_viweco = clean_num(s_bgvan.cell_value(15, 46)) or 1600000.0
        p_viweco_buy = clean_num(s_bgvan.cell_value(11, 46)) or 1312500.0
        p_belimo_5n = clean_num(s_bgvan.cell_value(15, 47)) or 2200000.0
        p_belimo_5n_buy = clean_num(s_bgvan.cell_value(11, 47)) or 1650000.0
        p_belimo_10n = clean_num(s_bgvan.cell_value(15, 48)) or 2800000.0
        p_belimo_10n_buy = clean_num(s_bgvan.cell_value(11, 48)) or 2098000.0
        p_belimo_20n = clean_num(s_bgvan.cell_value(15, 49)) or 3150000.0
        p_belimo_20n_buy = clean_num(s_bgvan.cell_value(11, 49)) or 2472500.0
        p_spring_box = clean_num(s_bgvan.cell_value(30, 51)) or clean_num(s_bgvan.cell_value(15, 51)) or 60000.0
        p_silicon_bolt = clean_num(s_bgvan.cell_value(30, 52)) or 40000.0
        p_fuse_handle_1 = clean_num(s_bgvan.cell_value(30, 38)) or 150000.0
        p_fuse_handle_2 = clean_num(s_bgvan.cell_value(42, 38)) or 225000.0
        p_spring_pair = clean_num(s_bgvan.cell_value(15, 39)) or 4000.0

        master_records.extend([
            {"MAT_ID": "MAT-FIREBOARD-10MM-BODY", "Category": "VẬT_LIỆU_CÁCH_NHIỆT", "Name": "Tấm chống cháy dày 10mm bọc thân van FD/MFD", "Unit": "m2", "Price": p_fb10_body, "Source_Trace": "BG VAN !R16C43"},
            {"MAT_ID": "MAT-FIREBOARD-20MM-BLADE", "Category": "VẬT_LIỆU_CÁCH_NHIỆT", "Name": "Tấm chống cháy dày 20mm làm lõi cánh van FD/MFD", "Unit": "m2", "Price": p_fb20_blade, "Source_Trace": "BG VAN !R16C44"},
            {"MAT_ID": "MAT-FIREBOARD-10MM-BLADE", "Category": "VẬT_LIỆU_CÁCH_NHIỆT", "Name": "Tấm chống cháy dày 10mm ốp 2 mặt cánh van FD/MFD", "Unit": "m2", "Price": p_fb10_blade, "Source_Trace": "BG VAN !R16C45"},
            {"MAT_ID": "MAT-INSULATION-EI-DAMPER", "Category": "VẬT_LIỆU_CÁCH_NHIỆT", "Name": "Bọc cách nhiệt chống cháy tấm MGO/Thạch cao 10mm (Tiêu chuẩn EI)", "Unit": "m2", "Price": p_fb10_body, "Source_Trace": "BG VAN !R16C43"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-FSLF24", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ chặn lửa ON/OFF 3.5Nm, 230V/24V (BELIMO FSLF24 / FSLF230)", "Unit": "Cái", "Price": p_fslf, "Purchase_Price": p_fslf_buy, "Source_Trace": "BG VAN !R16C41"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-BFN24", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ chặn lửa ON/OFF 9Nm, 230V/24V (BELIMO BFN24 / BFN230)", "Unit": "Cái", "Price": p_bfn, "Purchase_Price": p_bfn_buy, "Source_Trace": "BG VAN !R16C42"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-BF24", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ chặn lửa ON/OFF 18Nm, 24V (BELIMO BF 24 hoặc tương đương)", "Unit": "Cái", "Price": 5500000.0, "Source_Trace": "Bang_Tinh_Gia_Swind_Chuan_2026_Estimated"},
            {"MAT_ID": "MAT-ACTUATOR-HONEYWELL", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ chặn lửa ON/OFF (Honeywell)", "Unit": "Cái", "Price": p_honeywell, "Purchase_Price": p_honeywell_buy, "Source_Trace": "BG VAN !R16C46"},
            {"MAT_ID": "MAT-ACTUATOR-SWIND-VIWECO", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ chặn lửa ON/OFF Swind (Viweco V6006FS500S)", "Unit": "Cái", "Price": p_viweco, "Purchase_Price": p_viweco_buy, "Source_Trace": "BG VAN !R16C47"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-5NM", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ van gió Belimo ON/OFF 5Nm", "Unit": "Cái", "Price": p_belimo_5n, "Purchase_Price": p_belimo_5n_buy, "Source_Trace": "BG VAN !R16C48"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-10NM", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ van gió Belimo ON/OFF 10Nm", "Unit": "Cái", "Price": p_belimo_10n, "Purchase_Price": p_belimo_10n_buy, "Source_Trace": "BG VAN !R16C49"},
            {"MAT_ID": "MAT-ACTUATOR-BELIMO-20NM", "Category": "ĐỘNG_CƠ_VÀ_CẢM_BIẾN", "Name": "Động cơ van gió Belimo ON/OFF 20Nm", "Unit": "Cái", "Price": p_belimo_20n, "Purchase_Price": p_belimo_20n_buy, "Source_Trace": "BG VAN !R16C50"},
            {"MAT_ID": "MAT-DAMPER-HOP-LO-XO", "Category": "PHỤ_KIỆN_VAN", "Name": "Hộp lò xo van chặn lửa FD/MFD", "Unit": "Bộ", "Price": p_spring_box, "Source_Trace": "BG VAN !R31C52"},
            {"MAT_ID": "MAT-DAMPER-SILICON-BULONG-VIT", "Category": "PHỤ_KIỆN_VAN", "Name": "Silicon chống cháy + bu lông, ốc vít van FD/MFD", "Unit": "Bộ", "Price": p_silicon_bolt, "Source_Trace": "BG VAN !R31C53"},
            {"MAT_ID": "MAT-DAMPER-TAY-GAT-CAU-CHI-1", "Category": "PHỤ_KIỆN_VAN", "Name": "Tay gạt cầu chì van FD (Loại 1-2 cánh)", "Unit": "Bộ", "Price": p_fuse_handle_1, "Source_Trace": "BG VAN !R31C39"},
            {"MAT_ID": "MAT-DAMPER-TAY-GAT-CAU-CHI-2", "Category": "PHỤ_KIỆN_VAN", "Name": "Tay gạt cầu chì van FD (Loại 3-5 cánh)", "Unit": "Bộ", "Price": p_fuse_handle_2, "Source_Trace": "BG VAN !R43C39"},
            {"MAT_ID": "MAT-DAMPER-LO-XO-DOI", "Category": "PHỤ_KIỆN_VAN", "Name": "Lò xo van chặn lửa / đôi", "Unit": "Đôi", "Price": p_spring_pair, "Source_Trace": "BG VAN !R16C40"},
            {"MAT_ID": "ECCENTRIC_SHAFT", "Category": "PHỤ_KIỆN_VAN", "Name": "Cơ cấu lệch tâm / Lò xo đóng mở (NRD)", "Unit": "Bộ", "Price": 45000.0, "Source_Trace": "BG VAN CU"}
        ])

    # Khử trùng lặp theo MAT_ID (giữ nguyên thứ tự xuất hiện đầu tiên)
    deduped = []
    seen_ids = set()
    for rec in master_records:
        mid = rec["MAT_ID"]
        if mid not in seen_ids:
            seen_ids.add(mid)
            deduped.append(rec)

    return deduped

def save_master_data(data):
    os.makedirs(MASTER_DIR, exist_ok=True)
    # Split data into tech and price
    tech_data = []
    price_data = []
    
    for item in data:
        tech_item = {
            "MAT_ID": item.get("MAT_ID"),
            "Category": item.get("Category"),
            "Name": item.get("Name"),
            "Unit": item.get("Unit"),
            "Source_Trace": item.get("Source_Trace")
        }
        if "Weight_kg_m" in item:
            tech_item["Weight_kg_m"] = item["Weight_kg_m"]
        if "Density_kg_m2" in item:
            tech_item["Density_kg_m2"] = item["Density_kg_m2"]
        tech_data.append(tech_item)
        
        price_item = {
            "MAT_ID": item.get("MAT_ID"),
            "Price": item.get("Price", 0)
        }
        if "Price_m2" in item:
            price_item["Price_m2"] = item["Price_m2"]
        if "Purchase_Price" in item:
            price_item["Purchase_Price"] = item["Purchase_Price"]
        price_data.append(price_item)

    # 1. Save JSON
    with open(OUTPUT_JSON_TECH, 'w', encoding='utf-8') as f:
        json.dump(tech_data, f, ensure_ascii=False, indent=4)
        
    with open(OUTPUT_JSON_PRICE, 'w', encoding='utf-8') as f:
        json.dump(price_data, f, ensure_ascii=False, indent=4)

    prices_dir = os.path.join(MASTER_DIR, "prices")
    if os.path.exists(prices_dir):
        import glob
        price_files = sorted(glob.glob(os.path.join(prices_dir, "master_database_price_*.json")))
        if price_files:
            with open(price_files[-1], 'w', encoding='utf-8') as f:
                json.dump(price_data, f, ensure_ascii=False, indent=4)
        
    # 2. Save Excel Tech
    wb_tech = openpyxl.Workbook()
    ws_tech = wb_tech.active
    ws_tech.title = "05_DANH_MUC_KY_THUAT_GOC"
    
    headers_tech = ["STT", "MAT_ID", "PHÂN LOẠI", "TÊN VẬT TƯ / QUY CÁCH", "ĐVT", "TRỌNG LƯỢNG / BAREM", "NGUỒN DỮ LIỆU ĐỐI SOÁT"]
    ws_tech.append(headers_tech)
    
    # Styling
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="002060") # Swind Navy
    thin_border = Border(left=Side(style='thin', color='DDDDDD'),
                         right=Side(style='thin', color='DDDDDD'),
                         top=Side(style='thin', color='DDDDDD'),
                         bottom=Side(style='thin', color='DDDDDD'))
                         
    for col in range(1, len(headers_tech) + 1):
        cell = ws_tech.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for idx, item in enumerate(tech_data, 1):
        barem = item.get("Weight_kg_m") or item.get("Density_kg_m2") or ""
        barem_str = f"{barem} kg/m" if item.get("Weight_kg_m") else f"{barem} kg/m2" if item.get("Density_kg_m2") else ""
        
        row_vals = [
            idx,
            item.get("MAT_ID", ""),
            item.get("Category", ""),
            item.get("Name", ""),
            item.get("Unit", ""),
            barem_str,
            item.get("Source_Trace", "")
        ]
        ws_tech.append(row_vals)
        r = ws_tech.max_row
        
        for c in range(1, len(headers_tech) + 1):
            cell = ws_tech.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=9)
            cell.border = thin_border
            if c in [1, 2, 3, 5, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths tech
    ws_tech.column_dimensions['A'].width = 6
    ws_tech.column_dimensions['B'].width = 30
    ws_tech.column_dimensions['C'].width = 22
    ws_tech.column_dimensions['D'].width = 45
    ws_tech.column_dimensions['E'].width = 10
    ws_tech.column_dimensions['F'].width = 20
    ws_tech.column_dimensions['G'].width = 25
    
    wb_tech.save(OUTPUT_EXCEL_TECH)
    
    # 3. Save Excel Price
    wb_price = openpyxl.Workbook()
    ws_price = wb_price.active
    ws_price.title = "06_BANG_GIA_VAT_TU_AP_DUNG"
    
    headers_price = ["STT", "MAT_ID", "ĐƠN GIÁ GỐC (VNĐ)"]
    ws_price.append(headers_price)
    
    for col in range(1, len(headers_price) + 1):
        cell = ws_price.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    for idx, item in enumerate(price_data, 1):
        row_vals = [
            idx,
            item.get("MAT_ID", ""),
            item.get("Price", 0)
        ]
        ws_price.append(row_vals)
        r = ws_price.max_row
        
        for c in range(1, len(headers_price) + 1):
            cell = ws_price.cell(row=r, column=c)
            cell.font = Font(name="Arial", size=9)
            cell.border = thin_border
            if c in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")

    ws_price.column_dimensions['A'].width = 6
    ws_price.column_dimensions['B'].width = 30
    ws_price.column_dimensions['C'].width = 25
    
    wb_price.save(OUTPUT_EXCEL_PRICE)
    
    print(f"Exported {len(data)} verified master items to {OUTPUT_JSON_TECH} and {OUTPUT_JSON_PRICE}.")

if __name__ == "__main__":
    records = extract_all_master_data()
    save_master_data(records)
