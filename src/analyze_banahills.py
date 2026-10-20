import openpyxl
import os

file_path = "Copy of BG01.2_SLAND_NAMVIET- CỬA GIÓ VAN GIÓ TIÊU ÂM_BANA HILL H1- 08062026 VER4- van đa cánh BFN.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print(f"Sheets: {wb.sheetnames}")
for sheet_name in wb.sheetnames[:3]: # check first 3 sheets
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ---")
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(row): # Print non-empty rows
            print(f"Row {row_idx}: {row[:15]}")
        if row_idx > 25:
            break
