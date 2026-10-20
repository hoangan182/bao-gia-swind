import openpyxl
import glob

file_path = glob.glob("data/projects/Copy of BG01.2_SLAND_NAMVIET*.xlsx")[0]
print(f"Reading {file_path}...")
wb = openpyxl.load_workbook(file_path, data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for r in range(1, ws.max_row + 1):
        # Scan row for 'MFD'
        row_values = []
        has_mfd = False
        for c in range(1, min(ws.max_column + 1, 20)):
            val = ws.cell(r, c).value
            if val is not None:
                row_values.append(str(val))
                if "MFD" in str(val).upper():
                    has_mfd = True
        
        if has_mfd:
            print(f"Sheet '{sheet_name}', Row {r}: {' | '.join(row_values)}")

